"""Report files built from the tracker database.

Each builder returns (file name, bytes), or None when the selection has no rows.
`supplier` is a supplier key (core.supplier_key); titles and file names use core.supplier_name.
"""
import zipfile
from datetime import date, datetime
from io import BytesIO
from itertools import groupby
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import core

KINDS = (("STD", "Standard"), ("CUS", "Custom"))
TRACKER_COLS = [c for c in core.COLUMNS if c != "note"] + ["pacd", "matalan_rebate"]  # sheet columns A-W
RAW_COLS = TRACKER_COLS + ["note"]  # A-X
RAW_WIDTHS = [15.6, 13.4, 18.4, 23.1, 23.1, 28.2, 23.4, 9.0, 10.1, 20.5, 7.9, 14.8,
              11.0, 10.4, 10.8, 6.6, 6.6, 9.5, 12.0, 18.5, 11.8, 12.6, 20.0, 11.2]
XLSX_FORMATS = {
    "po_received_date": "mm-dd-yy",
    "shipment_date": "mm-dd-yy",
    **dict.fromkeys(["length_mm", "width_mm", "height_mm", "l_mm", "w_mm"], "#,##0"),
    "area_sqm": "#,##0.000",
    **dict.fromkeys(["unit_price", "final_price", "total_amount", "pacd", "matalan_rebate"], "$#,##0.00"),
}
HEAD_FONT = Font(bold=True, color="1F497D")
HEAD_FILL = PatternFill("solid", fgColor="B4C6E7")
BANNER_FONT = Font(bold=True, color="FFFFFF")
BANNER_FILL = PatternFill("solid", fgColor="4472C4")

SUMMARY_PREFIX = "Matalan Carton Shipment Tracking_Summary_"
SUMMARY_TITLE = "Carton Shipment Summary"
SUMMARY_HEADERS = ["Quarter", "Carton Type", "PACD Rate", "Rebate Rate", "Carton QTY", "Total Amount", "PACD",
                   "Matalan Rebate"]
SUMMARY_WIDTHS = [13, 13, 10, 10, 12, 15, 13, 16]
SUMMARY_MARGINS = {"left": 0.4, "right": 0.4, "top": 0.5, "bottom": 0.6, "header": 0.3, "footer": 0.3}  # inches
QTY_FORMAT = '#,##0;-#,##0;"-"'
MONEY_FORMAT = '#,##0.00;-#,##0.00;"-"'
PCT_FORMAT = "0.0%"
NAVY_HEX, BLUE_HEX, GREY_HEX = "1F497D", "4472C4", "7F7F7F"


def _font(size=10, bold=False, italic=False, color="000000"):
    return Font(name="Calibri", size=size, bold=bold, italic=italic, color=color)


BODY, BODY_BOLD, WHITE_BOLD = _font(), _font(bold=True), _font(bold=True, color="FFFFFF")
HEADING = _font(12, bold=True, color=NAVY_HEX)
NAVY_FILL = PatternFill("solid", fgColor=NAVY_HEX)
SUBTOTAL_FILL = PatternFill("solid", fgColor="DCE6F1")
GRAND_FILL = PatternFill("solid", fgColor="B4C6E7")
ROW_LINE = Border(bottom=Side("thin", color="D9D9D9"))
SUBTOTAL_LINE = Border(top=Side("thin", color=BLUE_HEX), bottom=Side("thin", color=BLUE_HEX))
GRAND_LINE = Border(top=Side("thin", color=NAVY_HEX), bottom=Side("double", color=NAVY_HEX))
HEADING_LINE = Border(bottom=Side("medium", color=BLUE_HEX))
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
HEADERS_ALIGN = [Alignment(horizontal=h, vertical="center", wrap_text=True, indent=1 if h == "left" else 0)
                 for h in ["left", "left", "center", "center", "right", "right", "right", "right"]]  # as the data
NOTE =Alignment(horizontal="left", vertical="top", wrap_text=True)


def safe_name(s):
    """Text usable in a Windows file or folder name."""
    for ch in '\\/:*?"<>|':
        s = s.replace(ch, "")
    return s.strip().rstrip(". ") or "Unknown"


def _headers(rows, cols):
    pacd, rebate = core.rate_headers(rows)
    return [{"pacd": pacd, "matalan_rebate": rebate}.get(c) or core.HEADERS[c] for c in cols]


def _bytes(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _banner(ws, row, text, width):
    cell = ws.cell(row, 1, text)
    cell.font, cell.fill, cell.alignment = BANNER_FONT, BANNER_FILL, Alignment(horizontal="left")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)


def _write_table(ws, row, rows, cols):
    """Header at `row`, data rows below it. Returns the first free row."""
    for i, h in enumerate(_headers(rows, cols), 1):
        cell = ws.cell(row, i, h)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, values in enumerate(rows, row + 1):
        for i, col in enumerate(cols, 1):
            value = values[col]
            if col in core.DATE_COLS and value:
                value = date.fromisoformat(value)
            elif col == "ebiz_so" and value and value.isdecimal():
                value = int(value)
            cell = ws.cell(r, i, value)
            if col in XLSX_FORMATS:
                cell.number_format = XLSX_FORMATS[col]
    return row + len(rows) + 1


def _add(values, more):
    return [round(a + b, 2) for a, b in zip(values, more)]


def _summary_block(con, key, fiscal_year):
    """(supplier name, [(quarter, [(carton type, (pacd %, rebate %), values)], quarter total)], grand total).

    values are [qty, total amount, pacd, rebate], money rounded to cents so totals add up as printed.
    A % is "Mixed" when the quarter's rows don't share one rate, "-" when the type has no rows. None: no rows.
    """
    summary = core.summary_rows(con, fiscal_year, supplier=key)
    if not summary:
        return None
    lines, grand = [], [0, 0, 0, 0]
    for q in sorted({row["quarter"] for row in summary}):
        kinds, q_total = [], [0, 0, 0, 0]
        for row in (item for item in summary if item["quarter"] == q):
            values = [row["cartons"], row["total_amount"], row["pacd"], row["matalan_rebate"]]
            kinds.append((row["carton_type"], (row["pacd_rate"], row["rebate_rate"]), values))
            q_total = _add(q_total, values)
        lines.append((q, kinds, q_total))
        grand = _add(grand, q_total)
    grand[0] = int(grand[0])
    return core.supplier_name(con, key), lines, grand


class _Rows:
    """Adds rows top to bottom and puts a page break before a section that would not fit on the page."""

    def __init__(self, ws, page_points):
        self.ws, self.page_points, self.next, self.used = ws, page_points, 1, 0.0

    def add(self, height):
        """Number of a new row of `height` points."""
        row = self.next
        self.ws.row_dimensions[row].height = height
        self.next += 1
        self.used += height
        return row

    def keep_together(self, height):
        if self.used + height > self.page_points and self.next > 1:
            self.ws.row_breaks.append(Break(id=self.next - 1))
            self.used = 0.0


def _page_points(widths):
    """Row height (points) that fits on one A4 portrait page once Excel fits the columns to the page width."""
    content = sum(7 * w for w in widths) * 0.75  # column width -> pixels -> points
    printable_w = (8.27 - SUMMARY_MARGINS["left"] - SUMMARY_MARGINS["right"]) * 72
    printable_h = (11.69 - SUMMARY_MARGINS["top"] - SUMMARY_MARGINS["bottom"]) * 72
    return printable_h / min(1, printable_w / content) * 0.95  # 5% spare for Excel's rounding


def _cell(ws, row, col, value=None, font=BODY, fmt=None, align=None, fill=None, border=None):
    cell = ws.cell(row, col, value)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    return cell


def _span(ws, row, first, last, value, **style):
    """One value across columns first..last (styled before merging, so the borders carry over)."""
    _cell(ws, row, first, value, **style)
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)


def _figures(ws, row, values, font=BODY, fill=None, border=None):
    """Carton QTY, Total Amount, PACD and Matalan Rebate in columns E-H."""
    for col, (value, fmt) in enumerate(zip(values, (QTY_FORMAT, MONEY_FORMAT, MONEY_FORMAT, MONEY_FORMAT)), 5):
        _cell(ws, row, col, value, font=font, fmt=fmt, align=RIGHT, fill=fill, border=border)


def _logos(ws, widths, top_px, height_px):
    """PacD logo at the top left, Matalan logo at the top right of the columns."""
    col_px = [7 * w for w in widths]
    for name, right in (("logo_left.png", False), ("logo_right.png", True)):
        path = ASSETS / name
        if not path.exists():
            continue
        img = XLImage(str(path))
        width_px = round(img.width * height_px / img.height)
        img.width, img.height = width_px, height_px
        x, col = (sum(col_px) - width_px if right else 0), 0
        while col < len(col_px) - 1 and x >= col_px[col]:
            x -= col_px[col]
            col += 1
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=col, colOff=pixels_to_EMU(x), row=0, rowOff=pixels_to_EMU(top_px)),
            ext=XDRPositiveSize2D(pixels_to_EMU(width_px), pixels_to_EMU(height_px)))
        ws.add_image(img)


def _title_block(ws, rows, ncols, widths, title, subtitle, now):
    """Rows 1-3: title, subtitle and generated time, centred between the PacD and Matalan logos."""
    _span(ws, rows.add(30), 1, ncols, title, font=_font(18, bold=True, color=NAVY_HEX), align=CENTER)
    _span(ws, rows.add(18), 1, ncols, subtitle, font=_font(11, color=GREY_HEX), align=CENTER)
    _span(ws, rows.add(16), 1, ncols, f"Generated {now:%d-%b-%Y %H:%M}",
          font=_font(9, italic=True, color=GREY_HEX), align=CENTER)
    _logos(ws, widths, top_px=20, height_px=44)


def _print_setup(ws, ncols, last_row, footer_left, now):
    """A4 portrait, one page wide, centred; footer with `footer_left`, 'Page x of y' and the generated date."""
    ws.print_area = f"A1:{get_column_letter(ncols)}{last_row}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(**SUMMARY_MARGINS)
    for part, text in ((ws.oddFooter.left, footer_left), (ws.oddFooter.center, "Page &P of &N"),
                       (ws.oddFooter.right, f"Generated {now:%d-%b-%Y}")):
        part.text, part.size, part.color = text, 8, GREY_HEX


def summary_xlsx(con, fiscal_year):
    """Title, one overview line per packaging supplier, then each supplier's quarters by carton type.

    Laid out for A4 portrait, one page wide; a supplier section never splits across pages.
    """
    span = core.date_range(con, fiscal_year=fiscal_year)
    if not span:
        return None
    blocks = [b for b in (_summary_block(con, key, fiscal_year) for key in core.suppliers(con)) if b]
    first, last = span
    now = datetime.now()
    wb = Workbook()
    wb.properties.title = f"{SUMMARY_TITLE} FY{fiscal_year}"
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    _widths(ws, SUMMARY_WIDTHS)
    ncols = len(SUMMARY_HEADERS)
    rows = _Rows(ws, _page_points(SUMMARY_WIDTHS))

    period = f"{first:%b %Y}" if (first.year, first.month) == (last.year, last.month) \
        else f"{first:%b %Y} – {last:%b %Y}"
    _title_block(ws, rows, ncols, SUMMARY_WIDTHS, SUMMARY_TITLE, f"Fiscal year {fiscal_year}  ·  Shipments {period}",
                 now)

    rows.add(12)
    _span(ws, rows.add(22), 1, ncols, "Overview by packaging supplier", font=HEADING, align=LEFT, border=HEADING_LINE)
    row = rows.add(30)
    _span(ws, row, 1, 4, "Packaging Supplier", font=WHITE_BOLD, fill=NAVY_FILL, align=LEFT)
    for col, header in enumerate(SUMMARY_HEADERS[4:], 5):
        _cell(ws, row, col, header, font=WHITE_BOLD, fill=NAVY_FILL, align=HEADERS_ALIGN[col - 1])
    everyone = [0, 0, 0, 0]
    for name, _, grand in blocks:
        row = rows.add(16)
        _span(ws, row, 1, 4, name, align=LEFT, border=ROW_LINE)
        _figures(ws, row, grand, border=ROW_LINE)
        everyone = _add(everyone, grand)
    everyone[0] = int(everyone[0])
    row = rows.add(18)
    _span(ws, row, 1, 4, "All suppliers", font=BODY_BOLD, align=LEFT, fill=GRAND_FILL, border=GRAND_LINE)
    _figures(ws, row, everyone, BODY_BOLD, GRAND_FILL, GRAND_LINE)

    for name, lines, grand in blocks:
        rows.keep_together(12 + 22 + 30 + len(lines) * (2 * 16 + 17) + 18)
        rows.add(12)
        _span(ws, rows.add(22), 1, ncols, name, font=HEADING, align=LEFT, border=HEADING_LINE)
        row = rows.add(30)
        for col, header in enumerate(SUMMARY_HEADERS, 1):
            _cell(ws, row, col, header, font=WHITE_BOLD, fill=NAVY_FILL, align=HEADERS_ALIGN[col - 1])
        for q, kinds, q_total in lines:
            top = rows.next
            for label, pcts, values in kinds:
                row = rows.add(16)
                _cell(ws, row, 1, f"Q{q} ({core.QUARTER_MONTHS[q]})" if row == top else None, align=LEFT,
                      border=ROW_LINE)
                _cell(ws, row, 2, label, align=LEFT, border=ROW_LINE)
                for col, pct in zip((3, 4), pcts):
                    _cell(ws, row, col, pct, fmt=PCT_FORMAT, align=CENTER, border=ROW_LINE)
                _figures(ws, row, values, border=ROW_LINE)
            ws.merge_cells(start_row=top, start_column=1, end_row=row, end_column=1)
            row = rows.add(17)
            _span(ws, row, 1, 4, f"Q{q} Total", font=BODY_BOLD, align=LEFT, fill=SUBTOTAL_FILL, border=SUBTOTAL_LINE)
            _figures(ws, row, q_total, BODY_BOLD, SUBTOTAL_FILL, SUBTOTAL_LINE)
        row = rows.add(18)
        _span(ws, row, 1, 4, "Grand Total", font=BODY_BOLD, align=LEFT, fill=GRAND_FILL, border=GRAND_LINE)
        _figures(ws, row, grand, BODY_BOLD, GRAND_FILL, GRAND_LINE)

    quarters = ", ".join(f"Q{q} {months}" for q, months in core.QUARTER_MONTHS.items())
    notes = [f"Quarters: {quarters}. December counts toward the next fiscal year.",
             "Rates apply by shipment date. Matalan Rebate = Unit Price × Carton QTY × Rebate Rate; "
             "PACD = (Total Amount − Matalan Rebate) × PACD Rate."]
    rows.keep_together(12 + 16 + 14 + 26)
    rows.add(12)
    _span(ws, rows.add(16), 1, ncols, "Notes", font=_font(9, bold=True, color=GREY_HEX), align=NOTE)
    for note, height in zip(notes, (14, 26)):
        _span(ws, rows.add(height), 1, ncols, note, font=_font(9, color=GREY_HEX), align=NOTE)

    _print_setup(ws, ncols, rows.next - 1, f"{SUMMARY_TITLE}  ·  FY{fiscal_year}", now)
    return f"{SUMMARY_PREFIX}{first:%b}-{last:%b} {last:%y}.xlsx", _bytes(wb)


XREF_HEADERS = ["SO#", "Tracker QTY", "Tracker Rebate", "eBiz QTY", "eBiz Rebate", "QTY Difference", "Status"]
XREF_WIDTHS = [16, 13, 15, 13, 15, 15, 16]
XREF_FORMATS = [None, QTY_FORMAT, MONEY_FORMAT, QTY_FORMAT, MONEY_FORMAT, QTY_FORMAT, None]
XREF_ALIGNS = [LEFT, RIGHT, RIGHT, RIGHT, RIGHT, RIGHT, CENTER]
STATUS_COLORS = {"Not in eBiz": "C00000", "Tracker > eBiz": "C55A11", "Tracker < eBiz": "C55A11", "Match": "548235"}


def _order(text):
    """An all-digit order number as a number (no 'stored as text' mark in Excel), anything else as is."""
    return int(text) if text.isdecimal() else text


def _xref_line(ws, row, values, font=BODY, fill=None, border=ROW_LINE):
    """Values from column A, formatted as the Cross Reference columns."""
    for col, (value, fmt, align) in enumerate(zip(values, XREF_FORMATS, XREF_ALIGNS), 1):
        _cell(ws, row, col, value, font=font, fmt=fmt, align=align, fill=fill, border=border)


def _xref_header(ws, row, headers):
    for col, (header, align) in enumerate(zip(headers, XREF_ALIGNS), 1):
        _cell(ws, row, col, header, font=WHITE_BOLD, fill=NAVY_FILL, align=align)


def cross_reference_xlsx(con, supplier, xref):
    """Tracker vs eBiz quantity and rebate per SO#, and on a second sheet the eBiz orders missing from the
    tracker, for one supplier.

    `xref` is core.cross_reference(con, supplier). Laid out like the Summary: A4 portrait, one page wide.
    """
    name = core.supplier_name(con, supplier)
    now = datetime.now()
    wb = Workbook()
    wb.properties.title = f"Cross Reference {name}"
    ws = wb.active
    ws.title = "Cross Reference"
    ws.sheet_view.showGridLines = False
    _widths(ws, XREF_WIDTHS)
    ncols = len(XREF_HEADERS)
    rows = _Rows(ws, _page_points(XREF_WIDTHS))
    _title_block(ws, rows, ncols, XREF_WIDTHS, "Cross Reference", f"{name}  ·  all tracker rows vs eBiz data", now)

    rows.add(12)
    _span(ws, rows.add(22), 1, ncols, "Tracker orders", font=HEADING, align=LEFT, border=HEADING_LINE)
    header_row = rows.add(30)
    _xref_header(ws, header_row, XREF_HEADERS)
    ws.print_title_rows = f"{header_row}:{header_row}"
    total = [0, 0, 0, 0]
    for o in xref["orders"]:
        values = [o["tracker_qty"], round(o["tracker_rebate"], 2), o["ebiz_qty"], round(o["ebiz_rebate"], 2)]
        row = rows.add(16)
        _xref_line(ws, row, [_order(o["so"])] + values + [values[0] - values[2], o["status"]])
        ws.cell(row, 7).font = _font(bold=True, color=STATUS_COLORS[o["status"]])
        total = _add(total, values)
    no_so = [r for kind_rows in xref["rows"].values() for r in kind_rows if not r["so"]]
    if no_so:
        qty = sum(r["carton_qty"] or 0 for r in no_so)
        rebate = round(sum((r["pacd"] or 0) + (r["matalan_rebate"] or 0) for r in no_so), 2)
        _xref_line(ws, rows.add(16), ["(No SO#)", qty, rebate, None, None, qty, None])  # qty - 0: the column adds up
        total = _add(total, [qty, rebate, 0, 0])
    _xref_line(ws, rows.add(18), ["Total"] + total + [round(total[0] - total[2], 2), None],
               BODY_BOLD, GRAND_FILL, GRAND_LINE)

    groups = {}
    for e in xref["ebiz_only"]:
        group = groups.setdefault(e["key"] or "(no order number)", [0, 0])
        group[0] += e["qty_shipped"] or 0
        group[1] += e["total_rebate"] or 0
    _print_setup(ws, ncols, rows.next - 1, f"Cross Reference  ·  {name}", now)

    # Own sheet: print titles repeat on every page of a sheet, so the tracker header must not sit above this list.
    ws = wb.create_sheet("eBiz not in tracker")
    ws.sheet_view.showGridLines = False
    _widths(ws, XREF_WIDTHS)
    rows = _Rows(ws, _page_points(XREF_WIDTHS))
    _title_block(ws, rows, ncols, XREF_WIDTHS, "Cross Reference", f"{name}  ·  eBiz orders not in tracker", now)
    rows.add(12)
    _span(ws, rows.add(22), 1, ncols, "eBiz orders not in tracker", font=HEADING, align=LEFT, border=HEADING_LINE)
    header_row = rows.add(30)
    _xref_header(ws, header_row, ["Order", "eBiz QTY", "eBiz Rebate"])
    ws.print_title_rows = f"{header_row}:{header_row}"
    if groups:
        for order in sorted(groups, key=core.so_order):
            qty, rebate = groups[order]
            _xref_line(ws, rows.add(16), [_order(order), qty, round(rebate, 2)])
        _xref_line(ws, rows.add(18), ["Total", sum(g[0] for g in groups.values()),
                                      round(sum(g[1] for g in groups.values()), 2)], BODY_BOLD, GRAND_FILL, GRAND_LINE)
    else:
        _span(ws, rows.add(16), 1, 3, "None", font=_font(italic=True, color=GREY_HEX), align=LEFT)
    _print_setup(ws, ncols, rows.next - 1, f"Cross Reference  ·  {name}", now)
    return f"Cross Reference_{safe_name(name)}_{now:%d-%b-%Y}.xlsx", _bytes(wb)


def raw_data_xlsx(con, supplier, fiscal_year, quarter):
    """One quarter for one supplier: a Standard Carton block, then a Custom Carton block."""
    blocks = [(label, core.rows(con, kind, supplier=supplier, fiscal_year=fiscal_year, quarters=[quarter]))
              for kind, label in KINDS]
    if not any(rows for _, rows in blocks):
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = f"Q{quarter} {fiscal_year}"
    row = 1
    for label, rows in blocks:
        if rows:
            _banner(ws, row, f"{label} Carton", len(RAW_COLS))
            row = _write_table(ws, row + 1, rows, RAW_COLS) + 1
    _widths(ws, RAW_WIDTHS)
    name = core.supplier_name(con, supplier)
    return f"Tracker Raw Data_{safe_name(name)}_Q{quarter}_{fiscal_year}.xlsx", _bytes(wb)


def shipment_tracking_xlsx(con, supplier):
    """All dated rows of one supplier as TRACK_STD and TRACK_CUS sheets, values only."""
    span = core.date_range(con, supplier=supplier)
    if not span:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    for kind, sheet in core.SHEETS.items():
        ws = wb.create_sheet(sheet)
        _write_table(ws, 1, core.rows(con, kind, supplier=supplier), RAW_COLS)
        _widths(ws, RAW_WIDTHS)
        ws.freeze_panes = "A2"
    first, last = span
    short = safe_name(core.supplier_name(con, supplier).split()[0])
    return f"Matalan Carton Shipment Tracking_NPS_BD {first:%b-%y} to {last:%b-%y} ({short}).xlsx", _bytes(wb)


ASSETS = Path(__file__).with_name("assets")
PAGE = landscape(A4)
MARGIN = 0.3 * inch
NAVY, BLUE = colors.HexColor("#1F497D"), colors.HexColor("#4472C4")
HEAD_BG, QUARTER_BG = colors.HexColor("#B4C6E7"), colors.HexColor("#D9E1F2")
ALT_BG, SUBTOTAL_BG, TOTAL_BG = colors.HexColor("#EBF1FA"), colors.HexColor("#FFF2CC"), colors.HexColor("#DCE6F1")
CELL = ParagraphStyle("cell", fontName="Helvetica", fontSize=5.5, leading=6.5)
HEAD = ParagraphStyle("head", parent=CELL, fontName="Helvetica-Bold", textColor=NAVY, alignment=1)
SUM_COLS = ["carton_qty", "final_price", "total_amount", "pacd", "matalan_rebate"]  # summed like the VBA
BASE_WIDTHS = {
    "sr": 16, "po_received_date": 30, "shipment_date": 30, "matalan_po": 30, "supplier_name": 48,
    "factory_name": 48, "packaging_supplier": 48, "ebiz_customer_ref": 44, "ebiz_so": 30, "carton_qty": 26,
    "carton_style": 40, "flute": 26, "carton_selection": 30, "length_mm": 22, "width_mm": 22, "height_mm": 22,
    "l_mm": 22, "w_mm": 22, "area_sqm": 24, "unit_price": 28, "final_price": 30, "total_amount": 40,
    "pacd": 34, "matalan_rebate": 34,
}
PADDING = [("LEFTPADDING", (0, 0), (-1, -1), 1.5), ("RIGHTPADDING", (0, 0), (-1, -1), 1.5),
           ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]


def _fmt(col, value):
    """PDF cell content: a string for numbers and dates, a wrapping Paragraph for text."""
    if value is None:
        return ""
    if col in core.DATE_COLS:
        d = date.fromisoformat(value)
        return f"{d.month}/{d.day}/{d.year}"
    if col in ("carton_qty", "length_mm", "width_mm", "height_mm", "l_mm", "w_mm"):
        return f"{value:,.0f}"
    if col == "area_sqm":
        return f"{value:.3f}"
    if col == "unit_price":
        return f"{value:.2f}"
    if col in ("final_price", "total_amount", "pacd", "matalan_rebate"):
        return f"$ {value:,.2f}"
    return Paragraph(escape(str(value)), CELL)


def _total_line(text, rows, cols):
    line = [text] + [""] * len(cols)
    for col in SUM_COLS:
        if col in cols:
            line[cols.index(col) + 1] = _fmt(col, sum(r[col] or 0 for r in rows))
    return line


def _block(label, rows, cols, widths):
    """One carton type: banners, header, per-quarter rows with subtotals, block total."""
    qty_i = cols.index("carton_qty") + 1
    data = []
    style = PADDING + [
        ("FONT", (0, 0), (-1, -1), "Helvetica", 5.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 2), (-1, -1), 0.25, HEAD_BG),
        ("ALIGN", (qty_i, 3), (-1, -1), "RIGHT"),
        ("BACKGROUND", (0, 2), (-1, 2), HEAD_BG),
    ]

    def banner(text, background, size, color=colors.white):
        r = len(data)
        data.append([text] + [""] * len(cols))
        style.extend([("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), background),
                      ("TEXTCOLOR", (0, r), (-1, r), color), ("FONT", (0, r), (-1, r), "Helvetica-Bold", size)])

    def total(text, total_rows, background):
        r = len(data)
        data.append(_total_line(text, total_rows, cols))
        style.extend([("SPAN", (0, r), (qty_i - 1, r)), ("BACKGROUND", (0, r), (-1, r), background),
                      ("FONT", (0, r), (-1, r), "Helvetica-Bold", 5.5), ("ALIGN", (0, r), (-1, r), "RIGHT")])

    banner(f"Full Tracker Table - {label}", NAVY, 9)
    banner(f"{label} Carton", BLUE, 7)
    data.append([Paragraph("Sr No.", HEAD)] + [Paragraph(escape(h), HEAD) for h in _headers(rows, cols)])
    sr = 0
    for (year, quarter), q_rows in groupby(rows, key=lambda r: (r["fiscal_year"], r["quarter"])):
        q_rows = list(q_rows)
        q_label = core.quarter_label(year, quarter)
        banner(f"{q_label} Quarter", QUARTER_BG, 6, NAVY)
        for i, r in enumerate(q_rows):
            sr += 1
            data.append([str(sr)] + [_fmt(c, r[c]) for c in cols])
            if i % 2:
                style.append(("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), ALT_BG))
        total(f"{q_label} Subtotal", q_rows, SUBTOTAL_BG)
    total(f"{label} Carton Total", rows, TOTAL_BG)
    return Table(data, colWidths=widths, repeatRows=3, style=TableStyle(style))


class _NumberedCanvas(rl_canvas.Canvas):
    """Canvas that adds 'Page x of y' once the page count is known."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        for page in self._pages:
            self.__dict__.update(page)
            self.setFont("Helvetica", 7)
            self.drawRightString(PAGE[0] - MARGIN, 0.2 * inch, f"Page {self._pageNumber} of {len(self._pages)}")
            super().showPage()
        super().save()


def _title_bar(title, avail):
    return Table([[title]], colWidths=[avail], style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), NAVY), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("FONT", (0, 0), (-1, -1), "Helvetica-Bold", 14), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))


def _pdf(story, title):
    """Landscape A4 PDF bytes with both logos, the generated time and 'Page x of y' on every page."""
    generated = datetime.now().strftime("Generated: %d-%b-%Y %I:%M %p")

    def decorate(c, doc):
        c.saveState()
        for name, right in (("logo_left.png", False), ("logo_right.png", True)):
            path = ASSETS / name
            if path.exists():
                img = ImageReader(str(path))
                iw, ih = img.getSize()
                h = 0.45 * inch
                w = iw * h / ih
                c.drawImage(img, PAGE[0] - MARGIN - w if right else MARGIN, PAGE[1] - 0.3 * inch - h, w, h,
                            mask="auto")
        c.setFont("Helvetica", 7)
        c.drawString(MARGIN, 0.2 * inch, generated)
        c.restoreState()

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=PAGE, leftMargin=MARGIN, rightMargin=MARGIN,
                            topMargin=0.9 * inch, bottomMargin=0.45 * inch, title=title)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate, canvasmaker=_NumberedCanvas)
    return buf.getvalue()


def full_tracker_pdf(con, supplier, fiscal_year, quarters, brand):
    """Full Tracker Sheet for PacD (brand "PacD") or for Matalan (brand "Matalan": no PACD column)."""
    cols = TRACKER_COLS if brand == "PacD" else [c for c in TRACKER_COLS if c != "pacd"]
    blocks = [(label, core.rows(con, kind, supplier=supplier, fiscal_year=fiscal_year, quarters=quarters))
              for kind, label in KINDS]
    all_rows = [r for _, rows in blocks for r in rows]
    if not all_rows:
        return None
    avail = PAGE[0] - 2 * MARGIN
    base = [BASE_WIDTHS["sr"]] + [BASE_WIDTHS[c] for c in cols]
    widths = [w * avail / sum(base) for w in base]
    name = core.supplier_name(con, supplier)
    title = f"Full Tracker Report ({name})"
    story = [_title_bar(title, avail), Spacer(1, 6)]
    for label, rows in blocks:
        if rows:
            story += [_block(label, rows, cols, widths), Spacer(1, 8)]
    qty_i = cols.index("carton_qty") + 1
    story.append(Table([_total_line("Tracker Grand Total", all_rows, cols)], colWidths=widths,
                       style=TableStyle(PADDING + [
                           ("SPAN", (0, 0), (qty_i - 1, 0)), ("BACKGROUND", (0, 0), (-1, -1), NAVY),
                           ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                           ("FONT", (0, 0), (-1, -1), "Helvetica-Bold", 6), ("ALIGN", (0, 0), (-1, -1), "RIGHT")])))
    qs = sorted(quarters)
    span = f"Q{qs[0]}" if len(qs) == 1 else f"Q{qs[0]}-Q{qs[-1]}"
    return f"Full Tracker Sheet for {brand}_{safe_name(name)}_{span}_{fiscal_year}.pdf", _pdf(story, title)


YELLOW = colors.HexColor("#FFFF99")
NOTE_STYLE = ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=8, textColor=colors.HexColor("#595959"))
DISCREPANCY_SECTIONS = [  # (order status, or None for rows without an SO#; title; bar colour; eBiz QTY column)
    ("Not in eBiz", "Tracker Orders Not Found in eBiz", colors.HexColor("#C00000"), False),
    ("Tracker > eBiz", "Quantity Mismatch (Tracker > eBiz)", colors.HexColor("#C55A11"), True),
    ("Tracker < eBiz", "Quantity Mismatch (Tracker < eBiz)", colors.HexColor("#C55A11"), True),
    (None, "Tracker Rows With No SO#", colors.HexColor("#505050"), False),
]
EBIZ_PDF_COLS = ["vendor", "vendor_code", "product", "order_number", "qty_shipped", "net_price_1000",
                 "gross_price_1000", "total_rebate", "currency"]
EBIZ_PDF_WIDTHS = [16, 90, 36, 60, 60, 30, 34, 34, 34, 24]  # SL, then EBIZ_PDF_COLS


def _bar(text, background, avail):
    return Table([[text]], colWidths=[avail], style=TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), background), ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("FONT", (0, 0), (-1, -1), "Helvetica-Bold", 10), ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))


def _section(text, background, avail, table):
    """A section bar that never sits alone at a page foot: it moves to the next page unless there is room below it
    for the "No discrepancies found." note, or (`table`) for the heading and first rows of a table, which can
    run to many pages."""
    return [CondPageBreak((1.3 if table else 0.65) * inch), _bar(text, background, avail), Spacer(1, 3)]


def _grid(header, lines, widths, total, right_from, highlight=(), banner=None):
    """Bordered table: optional blue banner, header, lines (alternate shading, `highlight` line indexes pale
    yellow) and a total line whose label spans the columns before `right_from`. Columns from `right_from`
    align right."""
    data = []
    style = PADDING + [("FONT", (0, 0), (-1, -1), "Helvetica", 5.5), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    if banner:
        data.append([banner] + [""] * (len(widths) - 1))
        style += [("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 0), BLUE),
                  ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7)]
    head = len(data)
    data.append([Paragraph(escape(h), HEAD) for h in header])
    style += [("BACKGROUND", (0, head), (-1, head), HEAD_BG), ("GRID", (0, head), (-1, -1), 0.25, HEAD_BG),
              ("ALIGN", (right_from, head + 1), (-1, -1), "RIGHT")]
    for i, line in enumerate(lines):
        data.append(line)
        r = len(data) - 1
        if i in highlight:
            style.append(("BACKGROUND", (0, r), (-1, r), YELLOW))
        elif i % 2:
            style.append(("BACKGROUND", (0, r), (-1, r), ALT_BG))
    r = len(data)
    data.append(total)
    style += [("SPAN", (0, r), (right_from - 1, r)), ("BACKGROUND", (0, r), (-1, r), TOTAL_BG),
              ("FONT", (0, r), (-1, r), "Helvetica-Bold", 5.5), ("ALIGN", (0, r), (-1, r), "RIGHT")]
    return Table(data, colWidths=widths, repeatRows=head + 1, style=TableStyle(style))


def _ebiz_fmt(col, value):
    """PDF cell content for an eBiz column."""
    if value is None:
        return ""
    if col == "qty_shipped":
        return f"{value:,.0f}"
    if col in core.EBIZ_REAL:
        return f"{value:,.2f}"
    return Paragraph(escape(str(value)), CELL)


def _no_discrepancies(avail):
    return Table([[Paragraph("No discrepancies found.", NOTE_STYLE)]], colWidths=[avail])


def discrepancy_pdf(con, supplier, xref):
    """Tracker rows that disagree with eBiz, then eBiz rows missing from the tracker, for one supplier.

    `xref` is core.cross_reference(con, supplier). Tracker rows matched through SO_2, and eBiz rows with a blank
    SO_1 and a non-blank SO_2, are pale yellow.
    """
    name = core.supplier_name(con, supplier)
    title = f"Discrepancy Report ({name})"
    avail = PAGE[0] - 2 * MARGIN
    orders = {o["so"]: o for o in xref["orders"]}
    qty_i = TRACKER_COLS.index("carton_qty") + 1
    story = [_title_bar(title, avail), Spacer(1, 6)]
    for status, heading, color, with_ebiz in DISCREPANCY_SECTIONS:
        extra = 1 if with_ebiz else 0
        base = [BASE_WIDTHS["sr"]] + [BASE_WIDTHS[c] for c in TRACKER_COLS] + [36] * extra  # 'eBiz QTY' fits
        widths = [w * avail / sum(base) for w in base]
        blocks = []
        for kind, label in KINDS:
            rows = [r for r in xref["rows"][kind] if (orders[r["so"]]["status"] if r["so"] else None) == status]
            if not rows:
                continue
            header = ["SL"] + _headers(rows, TRACKER_COLS) + ["eBiz QTY"] * extra
            lines = [[str(i)] + [_fmt(c, r[c]) for c in TRACKER_COLS]
                     + ([f"{orders[r['so']]['ebiz_qty']:,.0f}"] if extra else []) for i, r in enumerate(rows, 1)]
            highlight = {i for i, r in enumerate(rows) if r["so"] and orders[r["so"]]["via_so2"]}
            total = _total_line("Total", rows, TRACKER_COLS) + [""] * extra
            blocks += [_grid(header, lines, widths, total, qty_i, highlight, f"{label} Carton"), Spacer(1, 6)]
        story += _section(heading, color, avail, bool(blocks)) + (blocks or [_no_discrepancies(avail), Spacer(1, 6)])
    ebiz = xref["ebiz_only"]
    story += _section("eBiz Orders Not Found in Tracker", NAVY, avail, bool(ebiz))
    if ebiz:
        header = ["SL"] + [core.EBIZ_HEADERS[c] for c in EBIZ_PDF_COLS]
        lines = [[str(i)] + [_ebiz_fmt(c, e[c]) for c in EBIZ_PDF_COLS] for i, e in enumerate(ebiz, 1)]
        highlight = {i for i, e in enumerate(ebiz) if not e["so_1"] and e["so_2"]}
        total = ["Total", "", "", "", "", _ebiz_fmt("qty_shipped", sum(e["qty_shipped"] or 0 for e in ebiz)),
                 "", "", _ebiz_fmt("total_rebate", sum(e["total_rebate"] or 0 for e in ebiz)), ""]
        widths = [w * avail / sum(EBIZ_PDF_WIDTHS) for w in EBIZ_PDF_WIDTHS]
        story.append(_grid(header, lines, widths, total, 5, highlight))
    else:
        story.append(_no_discrepancies(avail))
    return f"Discrepancy Report_{safe_name(name)}_{date.today():%d-%b-%Y}.pdf", _pdf(story, title)


def build_all(con, suppliers, fiscal_year, quarters, discrepancy=False):
    """Every report for the chosen suppliers as (relative path, bytes).

    The Summary comes first, then each supplier's files under a folder named after the supplier. Tracker
    reports need rows in the selection. With `discrepancy`, suppliers that have eBiz rows also get the
    Discrepancy Report and the Cross Reference, which compare all their tracker rows (selection ignored).
    """
    items = [summary_xlsx(con, fiscal_year)]
    with_ebiz = set(core.ebiz_vendors(con)) if discrepancy else set()
    for key in suppliers:
        files = []
        if core.date_range(con, supplier=key, fiscal_year=fiscal_year, quarters=quarters):
            files += [full_tracker_pdf(con, key, fiscal_year, quarters, brand) for brand in ("PacD", "Matalan")]
            files += [raw_data_xlsx(con, key, fiscal_year, q) for q in sorted(quarters)]
            files.append(shipment_tracking_xlsx(con, key))
        if key in with_ebiz:
            xref = core.cross_reference(con, key)
            files += [discrepancy_pdf(con, key, xref), cross_reference_xlsx(con, key, xref)]
        folder = safe_name(core.supplier_name(con, key))
        items += [(f"{folder}/{name}", data) for name, data in filter(None, files)]
    return [item for item in items if item]


def month_name(con, suppliers, fiscal_year, quarters):
    """Full name of the latest shipment month in the selection (the Reports sub-folder), or None."""
    span = core.date_range(con, suppliers=suppliers, fiscal_year=fiscal_year, quarters=quarters)
    return f"{span[1]:%B}" if span else None


def to_zip(items):
    """build_all() items as one zip file, keeping their folders."""
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for path, data in items:
            z.writestr(path, data)
    return buf.getvalue()


def save_all(con, suppliers, fiscal_year, quarters, reports_dir, discrepancy=False):
    """Write build_all() under <reports_dir>/<Month>/. Existing files are overwritten."""
    month = month_name(con, suppliers, fiscal_year, quarters)
    if not month:
        return []
    paths = []
    for rel, data in build_all(con, suppliers, fiscal_year, quarters, discrepancy):
        path = Path(reports_dir) / month / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        paths.append(path)
    return paths
