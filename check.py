"""Self-check for the carton tracker. Run: python -m pytest check.py -v"""
import re
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pypdf
import pytest

import core
import reports
import web_ui

HERE = Path(__file__).parent
DEMO = HERE / "samples" / "Matalan Carton Shipment Tracking_NPS_BD Aug-26.xlsx"
TW = "Transworld Packaging (BD) Ltd."
UG = "Uniglory Packaging Industries Ltd."
TW_KEY = "TRANSWORLD"
UG_KEY = "UNIGLORY"


def insert(con, kind, **values):
    cols = ", ".join(values)
    marks = ", ".join("?" * len(values))
    with con:
        con.execute(f"INSERT INTO {core.TABLES[kind]} ({cols}) VALUES ({marks})", list(values.values()))


def test_shared_chrome_reserves_header_space_and_renders_floating_chip(monkeypatch):
    rendered = []
    monkeypatch.setattr(web_ui.st, "html", rendered.append)

    web_ui.app_chrome()
    web_ui.header("Dashboard")

    assert "mc-topbrand-spacer" in rendered[1]
    assert '[data-testid="stHeader"]::after' in rendered[0]
    assert "data:image/png;base64" in rendered[0]
    assert "height: 68px" in web_ui.THEME
    assert ".mc-footer-chip { position: fixed !important" in web_ui.THEME
    assert "position:fixed!important;right:12px!important;bottom:10px!important" in rendered[1]
    assert f"&copy; {date.today().year} PACD" in rendered[1]
    assert "Developed by" in rendered[1]
    assert "https://github.com/RecurSID" in rendered[1]
    assert "class='mc-topbrand'" not in rendered[1]
    assert "class='mc-topbrand'" not in rendered[2]
    assert "mc-header" in rendered[2]
    assert len(rendered) == 3


def test_quarter_mapping(tmp_path):
    con = core.connect(tmp_path / "t.db")
    cases = {
        "2025-12-01": (2026, 1), "2026-03-31": (2026, 1), "2026-04-01": (2026, 2),
        "2026-08-31": (2026, 3), "2026-09-01": (2026, 4), "2026-11-30": (2026, 4),
        "2026-12-01": (2027, 1),
    }
    for d in cases:
        insert(con, "STD", shipment_date=d, carton_qty=1, unit_price=1.0, total_amount=1.0)
    got = {r["shipment_date"]: (r["fiscal_year"], r["quarter"]) for r in con.execute("SELECT * FROM v_std")}
    assert got == cases


def test_rate_follows_shipment_date(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "CUS", shipment_date="2026-06-30", carton_qty=100, unit_price=2.0, total_amount=200.0)
    insert(con, "CUS", shipment_date="2026-07-01", carton_qty=100, unit_price=2.0, total_amount=200.0)
    insert(con, "STD", shipment_date="2026-07-01", carton_qty=100, unit_price=2.0, total_amount=200.0)
    cus = {r["shipment_date"]: r for r in con.execute("SELECT * FROM v_cus")}
    assert cus["2026-06-30"]["matalan_rebate"] == pytest.approx(20.0)  # 10%
    assert cus["2026-06-30"]["pacd"] == pytest.approx(13.5)  # (200 - 20) * 7.5%
    assert cus["2026-07-01"]["matalan_rebate"] == pytest.approx(10.0)  # 5% from Q3 2026
    assert cus["2026-07-01"]["pacd"] == pytest.approx(14.25)  # (200 - 10) * 7.5%
    std = con.execute("SELECT * FROM v_std").fetchone()
    assert std["matalan_rebate"] == pytest.approx(10.0)  # 5%
    assert std["pacd"] == pytest.approx(7.6)  # (200 - 10) * 4%


def test_rebate_uses_unit_price_and_pacd_uses_total(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-08-01", carton_qty=100, unit_price=2.0, total_amount=999.0)
    r = con.execute("SELECT * FROM v_std").fetchone()
    assert r["matalan_rebate"] == pytest.approx(10.0)  # 2.0 * 100 * 5%
    assert r["pacd"] == pytest.approx((999.0 - 10.0) * 0.04)


def test_row_without_date_has_no_rate_or_quarter(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "CUS", shipment_date=None, carton_qty=5, unit_price=1.0, total_amount=5.0)
    r = con.execute("SELECT * FROM v_cus").fetchone()
    assert (r["pacd"], r["matalan_rebate"], r["fiscal_year"], r["quarter"]) == (None, None, None, None)


def test_connect_is_repeatable_and_seeds_once(tmp_path):
    core.connect(tmp_path / "t.db")
    con = core.connect(tmp_path / "t.db")
    assert con.execute("SELECT COUNT(*) FROM rates").fetchone()[0] == 4


def test_connect_does_not_rewrite_a_current_schema(tmp_path, monkeypatch):
    core.connect(tmp_path / "t.db")

    def fail():
        raise AssertionError("schema re-run on a current database")

    monkeypatch.setattr(core, "_schema", fail)
    core.connect(tmp_path / "t.db")


def test_schema_v5_upgrade_removes_assignment_tables_without_losing_report_data(tmp_path):
    path = tmp_path / "t.db"
    con = core.connect(path)
    insert(con, "STD", factory_name="Factory A", packaging_supplier="Pack A",
           supplier_key="PACK", carton_qty=10)
    with con:
        con.execute("CREATE TABLE app_meta (key TEXT PRIMARY KEY, value TEXT)")
        con.execute("CREATE TABLE packaging_suppliers (id INTEGER PRIMARY KEY, name TEXT)")
        con.execute("CREATE TABLE assignment_factories (id INTEGER PRIMARY KEY, name TEXT)")
        con.execute("CREATE TABLE factory_assignments (factory_id INTEGER, packaging_supplier_id INTEGER)")
        con.execute("PRAGMA user_version = 4")
    con.close()

    con = core.connect(path)
    assert con.execute("PRAGMA user_version").fetchone()[0] == 5
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 1
    tables = {row[0] for row in con.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert not {"app_meta", "packaging_suppliers", "assignment_factories",
                "factory_assignments"} & tables

@pytest.mark.parametrize("name, key", [
    ("M&U  Packaging Ltd.", "M&U"),
    ("Transworld  Packaging ( BD ) Ltd", "TRANSWORLD"),
    ("UNIGLORY PACKAGING IND LTD.", "UNIGLORY"),
    ("  ", None),
    (None, None),
])
def test_supplier_key(name, key):
    assert core.supplier_key(name) == key


def rec(ship="2026-08-03", sup="Pack A Ltd.", qty=100, price=1.0, **kw):
    """One tracker row for make_xlsx. Keys are core.COLUMNS names."""
    row = {
        "shipment_date": datetime.fromisoformat(ship) if ship else None,
        "packaging_supplier": sup,
        "carton_qty": qty,
        "unit_price": price,
        "final_price": price,
        "total_amount": qty * price if qty and price else None,
    }
    row.update(kw)
    return row


def make_xlsx(path, std=(), cus=(), drop=None, formula_total=False, names=("TRACK_STD", "TRACK_CUS"),
              header_row=2):
    """Write a minimal tracker workbook: headers in header_row, data right below."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in zip(names, (std, cus)):
        ws = wb.create_sheet(name)
        cols = [c for c in core.COLUMNS if c != drop]
        for i, c in enumerate(cols, 1):
            ws.cell(header_row, i, core.HEADERS[c])
        for r, row in enumerate(rows, header_row + 1):
            for i, c in enumerate(cols, 1):
                value = "=1+1" if c == "total_amount" and formula_total else row.get(c)
                ws.cell(r, i, value)
    wb.save(path)
    return path


def test_demo_import_counts_and_totals(tmp_path):
    con = core.connect(tmp_path / "t.db")
    data = core.read_workbook(DEMO)
    assert {k: len(v) for k, v in data.items()} == {"STD": 88, "CUS": 118}
    assert core.save(con, data, DEMO.name) == {"STD": 88, "CUS": 118}

    def sums(view):
        return tuple(con.execute(f"SELECT SUM(carton_qty), SUM(total_amount), SUM(pacd), "
                                 f"SUM(matalan_rebate) FROM {view}").fetchone())

    assert sums("v_std") == pytest.approx((28001, 33646.9152, 1278.5828, 1682.3458), abs=1e-3)
    assert sums("v_cus") == pytest.approx((41529, 37815.6551, 2694.3654, 1890.7828), abs=1e-3)

    core.save(con, core.read_workbook(DEMO), DEMO.name)  # importing again replaces, not duplicates
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 88
    assert con.execute("SELECT COUNT(*) FROM track_cus").fetchone()[0] == 118


def test_template_rows_are_skipped(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec(), {"carton_selection": "Custom", "flute": "BC-Flute"}])
    assert len(core.read_workbook(path)["STD"]) == 1


def test_values_are_cleaned(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec(supplier_name="Lariz Fashion Ltd ", ebiz_so=7124457,
                                                   ebiz_customer_ref="#N/A")])
    row = core.read_workbook(path)["STD"][0]
    assert row["supplier_name"] == "Lariz Fashion Ltd"
    assert row["ebiz_so"] == "7124457"
    assert row["ebiz_customer_ref"] is None
    assert row["shipment_date"] == "2026-08-03"
    assert row["carton_qty"] == 100


def test_sheet_names_match_any_case(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec()], cus=[rec()], names=("track_std", "Track_Cus"))
    assert {k: len(v) for k, v in core.read_workbook(path).items()} == {"STD": 1, "CUS": 1}


def test_missing_sheet_is_an_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "TRACK_STD"
    wb.save(tmp_path / "a.xlsx")
    with pytest.raises(core.ImportFileError, match="TRACK_CUS"):
        core.read_workbook(tmp_path / "a.xlsx")


def test_missing_header_is_an_error(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec()], drop="carton_qty")
    with pytest.raises(core.ImportFileError, match="Carton QTY"):
        core.read_workbook(path)


def test_headers_in_row_1_are_found(tmp_path):
    # exported per-supplier trackers have no title row: headers in row 1, data from row 2
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec(), rec(qty=7)], cus=[rec()], header_row=1)
    data = core.read_workbook(path)
    assert [(r["_row"], r["carton_qty"]) for r in data["STD"]] == [(2, 100), (3, 7)]
    assert [r["_row"] for r in data["CUS"]] == [2]


def test_sheet_without_header_row_is_an_error(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec()], drop="shipment_date")
    with pytest.raises(core.ImportFileError, match="TRACK_STD: header row not found"):
        core.read_workbook(path)


def test_file_without_cached_values_is_an_error(tmp_path):
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec()], formula_total=True)
    with pytest.raises(core.ImportFileError, match="no calculated values"):
        core.read_workbook(path)


TODAY = date(2026, 10, 7)


def date_check(con, tmp_path, *recs, name="a.xlsx"):
    """Read a STD-only workbook of recs and run the date check. Returns (files, reused, issues)."""
    files = {name: core.read_workbook(make_xlsx(tmp_path / name, std=recs))}
    reused, issues = core.check_dates(con, files, today=TODAY)
    return files, reused, issues


def proposals(issues):
    return [(i["Excel row"], i["Problem"], i["PO Received Date"], i["Shipment Date"]) for i in issues]


def test_dates_that_look_right_are_not_listed(tmp_path):
    con = core.connect(tmp_path / "t.db")
    _, reused, issues = date_check(con, tmp_path,
                                   rec(ship="2025-12-21", po_received_date=datetime(2025, 12, 6)),
                                   rec(ship="2026-04-21", po_received_date=datetime(2025, 12, 25)),  # long lead time
                                   rec(ship="2026-04-21"))  # no PO date
    assert (reused, issues) == (0, [])


def test_text_dates_are_read_day_first(tmp_path):
    con = core.connect(tmp_path / "t.db")
    files, _, issues = date_check(con, tmp_path, rec(ship="2025-12-21", po_received_date="18/12/2025"))
    assert files["a.xlsx"]["STD"][0]["po_received_date"] == "2025-12-18"
    assert proposals(issues) == [(3, "PO Received Date typed as text, read as day/month/year",
                                  date(2025, 12, 18), date(2025, 12, 21))]
    assert issues[0]["PO Received Date in file"] == "18/12/2025"


def test_swapped_po_date_gets_a_suggestion(tmp_path):
    # factory typed 06/12/2025 on a month-first PC: Excel stored 12 June
    con = core.connect(tmp_path / "t.db")
    _, _, issues = date_check(con, tmp_path, rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12)))
    assert proposals(issues) == [(3, "PO Received Date: day and month look swapped",
                                  date(2025, 12, 6), date(2025, 12, 21))]
    assert issues[0]["PO Received Date in file"] == "12-Jun-2025"


def test_swapped_shipment_date_gets_a_suggestion(tmp_path):
    con = core.connect(tmp_path / "t.db")
    _, _, issues = date_check(con, tmp_path, rec(ship="2026-02-04", po_received_date=datetime(2026, 3, 20)),
                              rec(ship="2026-12-05"))  # after TODAY
    assert proposals(issues) == [
        (3, "Shipment Date: day and month look swapped", date(2026, 3, 20), date(2026, 4, 2)),
        (4, "Shipment Date is in the future; day and month look swapped", None, date(2026, 5, 12))]


def test_po_date_without_its_year_gets_the_shipment_year(tmp_path):
    con = core.connect(tmp_path / "t.db")
    _, _, issues = date_check(con, tmp_path, rec(ship="2026-06-15", po_received_date=datetime(1900, 6, 13)),
                              rec(ship="2026-01-05", po_received_date=datetime(1900, 12, 28)))
    assert proposals(issues) == [
        (3, "PO Received Date is before 2020; the year looks wrong", date(2026, 6, 13), date(2026, 6, 15)),
        (4, "PO Received Date is before 2020; the year looks wrong", date(2025, 12, 28), date(2026, 1, 5))]


def test_po_after_shipment_is_listed_without_a_suggestion(tmp_path):
    con = core.connect(tmp_path / "t.db")
    _, _, issues = date_check(con, tmp_path, rec(ship="2026-04-07", po_received_date=datetime(2026, 4, 18)),
                              rec(ship="2026-04-07", po_received_date=datetime(1900, 6, 13)))
    assert proposals(issues) == [
        (3, "PO Received Date is after Shipment Date", date(2026, 4, 18), date(2026, 4, 7)),
        (4, "PO Received Date is before 2020", date(1900, 6, 13), date(2026, 4, 7))]


def test_missing_or_bad_shipment_date_is_listed_and_kept(tmp_path):
    con = core.connect(tmp_path / "t.db")
    files, _, issues = date_check(con, tmp_path, rec(), rec(ship=None, shipment_date="TBC"), rec(ship=None))
    assert proposals(issues) == [
        (4, "Shipment Date is not a date; the row is left out of reports", None, None),
        (5, "Shipment Date is empty; the row is left out of reports", None, None)]
    core.save_many(con, files)
    assert con.execute("SELECT COUNT(*) FROM v_std WHERE quarter IS NULL").fetchone()[0] == 2


def test_reviewed_dates_are_imported_and_reused_next_time(tmp_path):
    con = core.connect(tmp_path / "t.db")
    bad = rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12))
    files, _, issues = date_check(con, tmp_path, bad, rec())
    core.apply_dates(issues, [(date(2025, 12, 5), date(2025, 12, 20))])  # user's own choice
    core.save_many(con, files)
    assert [tuple(r) for r in con.execute("SELECT po_received_date, shipment_date FROM track_std ORDER BY id")] == [
        ("2025-12-05", "2025-12-20"), (None, "2026-08-03")]

    # next month's cumulative file carries the same row: corrected again, not listed
    files, reused, issues = date_check(con, tmp_path, bad, rec(), rec(qty=7), name="b.xlsx")
    assert (reused, issues) == (1, [])
    assert files["b.xlsx"]["STD"][0]["po_received_date"] == "2025-12-05"
    assert files["b.xlsx"]["STD"][0]["shipment_date"] == "2025-12-20"
    saved = core.saved_dates(con)
    assert len(saved) == 1 and saved[0].pop("key") and saved[0].pop("Saved at")
    assert saved[0] == {
        "Sheet": "TRACK_STD", "Packaging supplier": "Pack A Ltd.", "eBiz Sales Order #": None, "Carton QTY": 100,
        "PO Received Date in file": "12-Jun-2025", "Shipment Date in file": "21-Dec-2025",
        "PO Received Date": "05-Dec-2025", "Shipment Date": "20-Dec-2025"}


def test_changed_row_is_checked_again(tmp_path):
    con = core.connect(tmp_path / "t.db")
    files, _, issues = date_check(con, tmp_path, rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12)))
    core.apply_dates(issues, [(date(2025, 12, 6), date(2025, 12, 21))])
    core.save_many(con, files)
    _, reused, issues = date_check(con, tmp_path, rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12),
                                                      qty=99), name="b.xlsx")
    assert reused == 0 and len(issues) == 1


def test_forgotten_correction_is_checked_again(tmp_path):
    con = core.connect(tmp_path / "t.db")
    bad = rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12))
    files, _, issues = date_check(con, tmp_path, bad)
    core.apply_dates(issues, [(date(2025, 12, 6), date(2025, 12, 21))])
    core.save_many(con, files)
    core.forget_dates(con, [core.saved_dates(con)[0]["key"]])
    assert core.saved_dates(con) == []
    _, reused, issues = date_check(con, tmp_path, bad, name="b.xlsx")
    assert reused == 0 and len(issues) == 1


def test_old_database_gets_the_corrections_table(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-08-03", carton_qty=1, unit_price=1.0, total_amount=1.0)
    con.execute("DROP TABLE date_fixes")
    con.execute("PRAGMA user_version = 1")
    con.close()
    con = core.connect(tmp_path / "t.db")
    assert core.saved_dates(con) == []
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 1


def test_reimport_replaces_only_overlapping_supplier_months(tmp_path):
    con = core.connect(tmp_path / "t.db")
    first = make_xlsx(tmp_path / "a.xlsx", std=[rec(sup="A"), rec(sup="A"), rec(sup="B"),
                                                rec(ship="2026-07-15", sup="A")])
    core.save(con, core.read_workbook(first), "a.xlsx")
    second = core.read_workbook(make_xlsx(tmp_path / "b.xlsx", std=[rec(sup="A", qty=5)]))
    assert core.preview(con, second) == [{"Sheet": "TRACK_STD", "Supplier": "A", "Spellings in file": "A",
                                          "Shipment month": "2026-08", "Rows in file": 1,
                                          "Rows replaced in DB": 2}]
    core.save(con, second, "b.xlsx")
    got = sorted((r["packaging_supplier"], r["shipment_date"], r["carton_qty"])
                 for r in con.execute("SELECT * FROM track_std"))
    assert got == [("A", "2026-07-15", 100), ("A", "2026-08-03", 5), ("B", "2026-08-03", 100)]


def test_supplier_spellings_merge(tmp_path):
    con = core.connect(tmp_path / "t.db")
    spellings = ["M&U Packaging Ltd", "M&U Packaging Ltd", "M&U  Packaging Ltd.", "m&u packaging ltd."]
    data = core.read_workbook(make_xlsx(tmp_path / "a.xlsx", std=[rec(sup=s, qty=10 * i)
                                                                  for i, s in enumerate(spellings, 1)]))
    assert core.preview(con, data) == [{
        "Sheet": "TRACK_STD", "Supplier": "M&U",
        "Spellings in file": "M&U  Packaging Ltd. / M&U Packaging Ltd / m&u packaging ltd.",
        "Shipment month": "2026-08", "Rows in file": 4, "Rows replaced in DB": 0}]
    core.save(con, data, "a.xlsx")
    assert core.suppliers(con) == ["M&U"]
    assert core.supplier_name(con, "M&U") == "M&U Packaging Ltd"
    _, rows = summary_sheet(reports.summary_xlsx(con, 2026)[1])
    assert [r[0] for r in rows if r[0] and r[0].startswith("M&U")] == ["M&U Packaging Ltd"] * 2  # overview, block
    assert [r[4] for r in rows if r[0] == "Grand Total"] == [10 + 20 + 30 + 40]


def test_spelling_fixed_on_reimport_replaces_rows(tmp_path):
    con = core.connect(tmp_path / "t.db")
    old = make_xlsx(tmp_path / "a.xlsx", std=[rec(sup="Uniglory Packging Industries Ltd.")] * 2)
    core.save(con, core.read_workbook(old), "a.xlsx")
    new = make_xlsx(tmp_path / "b.xlsx", std=[rec(sup="Uniglory Packaging Industries Ltd.")] * 2)
    core.save(con, core.read_workbook(new), "b.xlsx")
    assert [r[0] for r in con.execute("SELECT packaging_supplier FROM track_std")] == [UG] * 2


def test_several_files_import_together(tmp_path):
    con = core.connect(tmp_path / "t.db")
    files = {
        "a.xlsx": core.read_workbook(make_xlsx(tmp_path / "a.xlsx", std=[rec(sup="Alpha Ltd")] * 2)),
        "b.xlsx": core.read_workbook(make_xlsx(tmp_path / "b.xlsx", std=[rec(sup="Beta Ltd")],
                                               cus=[rec(sup="Alpha Ltd")])),  # other sheet: no clash
    }
    assert core.clashes(files) == []
    assert core.save_many(con, files) == {"a.xlsx": {"STD": 2, "CUS": 0}, "b.xlsx": {"STD": 1, "CUS": 1}}
    got = sorted(tuple(r) for r in con.execute(
        "SELECT supplier_key, source_file FROM track_std UNION ALL SELECT supplier_key, source_file FROM track_cus"))
    assert got == [("ALPHA", "a.xlsx"), ("ALPHA", "a.xlsx"), ("ALPHA", "b.xlsx"), ("BETA", "b.xlsx")]


def test_files_with_the_same_supplier_month_clash(tmp_path):
    con = core.connect(tmp_path / "t.db")
    files = {
        "cumulative.xlsx": core.read_workbook(make_xlsx(tmp_path / "a.xlsx", std=[
            rec(ship="2026-07-10", sup="Alpha Ltd"), rec(sup="Alpha Ltd")])),
        "august.xlsx": core.read_workbook(make_xlsx(tmp_path / "b.xlsx", std=[rec(sup="ALPHA LTD.")])),
    }
    assert core.clashes(files) == [{"Sheet": "TRACK_STD", "Supplier": "ALPHA", "Shipment month": "2026-08",
                                    "Files": "cumulative.xlsx / august.xlsx"}]
    with pytest.raises(ValueError, match="ALPHA"):
        core.save_many(con, files)
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 0


def test_failed_batch_leaves_database_unchanged(tmp_path):
    con = core.connect(tmp_path / "t.db")
    good = core.read_workbook(make_xlsx(tmp_path / "a.xlsx", std=[rec(sup="Alpha Ltd")]))
    broken = {"STD": [{"supplier_key": "BETA", "shipment_date": "2026-08-03"}], "CUS": []}  # missing columns
    with pytest.raises(KeyError):
        core.save_many(con, {"a.xlsx": good, "b.xlsx": broken})
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 0


SEEDED = [
    {"carton_type": "STD", "effective_from": "2025-12-01", "pacd_pct": 0.04, "matalan_pct": 0.05},
    {"carton_type": "STD", "effective_from": "2026-07-01", "pacd_pct": 0.04, "matalan_pct": 0.05},
    {"carton_type": "CUS", "effective_from": "2025-12-01", "pacd_pct": 0.075, "matalan_pct": 0.10},
    {"carton_type": "CUS", "effective_from": "2026-07-01", "pacd_pct": 0.075, "matalan_pct": 0.05},
]


def test_seeded_rates(tmp_path):
    assert core.get_rates(core.connect(tmp_path / "t.db")) == SEEDED


def test_rate_change_applies_to_existing_rows(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "CUS", shipment_date="2026-09-10", carton_qty=100, unit_price=2.0, total_amount=200.0)
    rates = core.get_rates(con)
    rates[-1]["matalan_pct"] = 0.08
    core.save_rates(con, rates)
    assert con.execute("SELECT matalan_rebate FROM v_cus").fetchone()[0] == pytest.approx(16.0)


@pytest.mark.parametrize("change, message", [
    ({"carton_type": "XYZ"}, "STD or CUS"),
    ({"effective_from": "None"}, "must be a date"),
    ({"pacd_pct": 1.5}, "between 0 and 100"),
    ({"matalan_pct": float("nan")}, "between 0 and 100"),
    ({"effective_from": "2026-07-01"}, "duplicate"),
])
def test_invalid_rates_are_rejected(tmp_path, change, message):
    con = core.connect(tmp_path / "t.db")
    rates = core.get_rates(con)
    rates[0].update(change)  # rates[0] is STD 2025-12-01; the duplicate case collides with STD 2026-07-01
    with pytest.raises(ValueError, match=message):
        core.save_rates(con, rates)
    assert core.get_rates(con) == SEEDED


def test_empty_rates_are_rejected(tmp_path):
    con = core.connect(tmp_path / "t.db")
    with pytest.raises(ValueError, match="At least one"):
        core.save_rates(con, [])


def test_rates_that_leave_shipments_unrated_are_rejected(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-01-15", carton_qty=1, unit_price=1.0, total_amount=1.0)
    message = ("1 shipment(s) from 2026-01-15 on would have no rate. "
               "Add a rate that starts on or before 2026-01-15.")
    with pytest.raises(ValueError, match=re.escape(message)):
        core.save_rates(con, SEEDED[1:])  # drops STD 2025-12-01
    assert core.get_rates(con) == SEEDED


def test_unrated_count(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2025-06-01", carton_qty=1, unit_price=1.0, total_amount=1.0)
    insert(con, "CUS", shipment_date="2026-08-01", carton_qty=1, unit_price=1.0, total_amount=1.0)
    insert(con, "CUS", shipment_date=None, carton_qty=1, unit_price=1.0, total_amount=1.0)
    assert core.unrated_count(con) == 1


@pytest.fixture
def demo_con(tmp_path):
    con = core.connect(tmp_path / "demo.db")
    core.save(con, core.read_workbook(DEMO), DEMO.name)
    return con


def test_queries(demo_con):
    assert core.suppliers(demo_con) == [TW_KEY, UG_KEY]
    assert core.supplier_name(demo_con, TW_KEY) == TW
    assert core.supplier_name(demo_con, "NOBODY") == "NOBODY"
    assert core.fiscal_years(demo_con) == [2026]
    assert core.quarters(demo_con, supplier=TW_KEY, fiscal_year=2026) == [3]
    assert core.date_range(demo_con) == (date(2026, 8, 1), date(2026, 8, 31))
    assert core.date_range(demo_con, fiscal_year=2027) is None
    assert len(core.rows(demo_con, "STD", supplier=TW_KEY, fiscal_year=2026, quarters=[3])) == 55
    assert len(core.rows(demo_con, "CUS", supplier=TW_KEY)) == 91
    dates = [r["shipment_date"] for r in core.rows(demo_con, "CUS", supplier=UG_KEY)]
    assert dates == sorted(dates)
    qty, total, pacd, rebate = core.quarter_totals(demo_con, "STD", TW_KEY, 2026)[3]
    assert qty == sum(r["carton_qty"] for r in core.rows(demo_con, "STD", supplier=TW_KEY))
    assert core.quarter_label(2026, 1) == "Q1 2026 (Dec-Mar)"


def test_tracker_rows_can_be_manually_edited(demo_con):
    row = core.editable_rows(demo_con, "STD", TW_KEY)[0]
    row["po_received_date"] = "04/08/2026"
    row["shipment_date"] = "2026-08-09"
    row["matalan_po"] = "PO-CORRECTED"
    row["carton_qty"] = "321"
    assert core.update_tracker_rows(demo_con, "STD", [row]) == 1
    saved = demo_con.execute(
        "SELECT po_received_date, shipment_date, matalan_po, carton_qty "
        "FROM track_std WHERE id = ?", (row["id"],)).fetchone()
    assert tuple(saved) == ("2026-08-04", "2026-08-09", "PO-CORRECTED", 321)


def test_invalid_tracker_edit_rolls_back(demo_con):
    rows = core.editable_rows(demo_con, "STD", TW_KEY)[:2]
    before = [r["matalan_po"] for r in rows]
    rows[0]["matalan_po"] = "SHOULD-NOT-SAVE"
    rows[1]["shipment_date"] = "not a date"
    with pytest.raises(ValueError, match="Shipment Date"):
        core.update_tracker_rows(demo_con, "STD", rows)
    after = [r[0] for r in demo_con.execute(
        "SELECT matalan_po FROM track_std WHERE id IN (?, ?) ORDER BY id",
        [r["id"] for r in rows])]
    assert after == before


def test_live_summary_rows_match_report_totals(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-08-01", packaging_supplier="Pack Ltd",
           supplier_key="PACK", supplier_name="Supplier A", factory_name="Factory A",
           carton_qty=10, unit_price=2.0, total_amount=20.0)
    insert(con, "CUS", shipment_date="2026-08-02", packaging_supplier="Pack Ltd",
           supplier_key="PACK", supplier_name="Supplier B", factory_name="Factory B",
           carton_qty=5, unit_price=4.0, total_amount=20.0)
    rows = core.summary_rows(con, 2026)
    assert [(row["carton_type"], row["cartons"], row["total_amount"])
            for row in rows] == [("Standard", 10, 20.0), ("Custom", 5, 20.0)]
    assert all(row["quarter_label"] == "Q3 (Jul-Aug)" for row in rows)
    assert rows[0]["pacd_rate"] == 0.04 and rows[1]["pacd_rate"] == 0.075
    assert core.summary_rows(con, 2027) == []


def test_rate_headers():
    one = [{"pacd_pct": 0.075, "matalan_pct": 0.05}] * 2
    mixed = [{"pacd_pct": 0.075, "matalan_pct": 0.05}, {"pacd_pct": 0.075, "matalan_pct": 0.10}]
    unrated = [{"pacd_pct": 0.04, "matalan_pct": 0.05}, {"pacd_pct": None, "matalan_pct": None}]
    assert core.rate_headers(one) == ("7.5% for PACD", "5% Rebates for Matalan")
    assert core.rate_headers(mixed) == ("7.5% for PACD", "Rebates for Matalan")
    assert core.rate_headers(unrated) == ("PACD", "Rebates for Matalan")
    assert core.rate_headers([]) == ("PACD", "Rebates for Matalan")


def test_safe_name():
    assert reports.safe_name('A/B: "C". ') == "AB C"
    assert reports.safe_name(TW) == "Transworld Packaging (BD) Ltd"


def summary_sheet(data):
    ws = openpyxl.load_workbook(BytesIO(data)).active
    return ws, [[c.value for c in r] for r in ws.iter_rows()]


def test_summary_xlsx(demo_con):
    name, data = reports.summary_xlsx(demo_con, 2026)
    assert name == "Matalan Carton Shipment Tracking_Summary_Aug-Aug 26.xlsx"
    ws, rows = summary_sheet(data)
    assert rows[0][0] == "Carton Shipment Summary"
    assert rows[1][0] == "Fiscal year 2026  ·  Shipments Aug 2026"
    assert len(ws._images) == 2  # PacD and Matalan logos

    overview = [r for r in rows if r[0] in (TW, UG, "All suppliers") and r[4] is not None]
    assert [r[0] for r in overview] == [TW, UG, "All suppliers"]
    assert overview[2][4] == 28001 + 41529
    assert overview[2][5] == pytest.approx(33646.9152 + 37815.6551, abs=0.05)
    assert overview[2][7] == pytest.approx(1682.3458 + 1890.7828, abs=0.05)

    head = next(i for i, r in enumerate(rows) if r[0] == TW and r[4] is None)
    assert rows[head + 1] == ["Quarter", "Carton Type", "PACD Rate", "Rebate Rate", "Carton QTY", "Total Amount",
                              "PACD", "Matalan Rebate"]
    std, cus, total = rows[head + 2:head + 5]
    assert std[:4] == ["Q3 (Jul-Aug)", "Standard", 0.04, 0.05]
    assert cus[:4] == [None, "Custom", 0.075, 0.05]  # quarter cell merged over both lines
    assert total[0] == "Q3 Total"
    assert total[5] == pytest.approx(std[5] + cus[5], abs=1e-9)  # Total adds the rounded values
    grand = [r for r in rows if r[0] == "Grand Total"]
    assert len(grand) == 2 and sum(r[4] for r in grand) == overview[2][4]
    assert sum(r[7] for r in grand) == pytest.approx(overview[2][7], abs=1e-9)

    assert ws.page_setup.orientation == "portrait" and ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage and not ws.sheet_view.showGridLines
    assert ws.oddFooter.center.text == "Page &P of &N"
    assert reports.summary_xlsx(demo_con, 2027) is None


def test_summary_shows_mixed_and_missing_rates(tmp_path):
    con = core.connect(tmp_path / "t.db")
    for d in ("2026-07-10", "2026-08-10"):
        insert(con, "STD", shipment_date=d, packaging_supplier="A", supplier_key="A", carton_qty=10,
               unit_price=1.0, total_amount=10.0)
    core.save_rates(con, core.get_rates(con) + [
        {"carton_type": "STD", "effective_from": "2026-08-01", "pacd_pct": 0.03, "matalan_pct": 0.05}])
    assert core.quarter_rates(con, "STD", "A", 2026) == {3: (None, 0.05)}
    _, rows = summary_sheet(reports.summary_xlsx(con, 2026)[1])
    head = next(i for i, r in enumerate(rows) if r[0] == "A" and r[4] is None)
    assert rows[head + 2][1:4] == ["Standard", "Mixed", 0.05]
    assert rows[head + 3][1:4] == ["Custom", "-", "-"]  # no Custom shipments


def test_raw_data_xlsx(demo_con):
    name, data = reports.raw_data_xlsx(demo_con, TW_KEY, 2026, 3)
    assert name == "Tracker Raw Data_Transworld Packaging (BD) Ltd_Q3_2026.xlsx"
    ws = openpyxl.load_workbook(BytesIO(data)).active
    assert ws.title == "Q3 2026"
    assert ws["A1"].value == "Standard Carton"
    assert [c.value for c in ws[2]][21:24] == ["4% for PACD", "5% Rebates for Matalan", "NOTE"]
    assert ws["A59"].value == "Custom Carton"  # 55 STD rows on 3-57, blank row 58
    assert [c.value for c in ws[60]][21:23] == ["7.5% for PACD", "5% Rebates for Matalan"]
    assert ws.max_row == 151  # 91 CUS rows on 61-151
    assert isinstance(ws["B3"].value, datetime)
    assert isinstance(ws["H3"].value, int)  # eBiz Sales Order #
    assert reports.raw_data_xlsx(demo_con, TW_KEY, 2026, 1) is None


def test_shipment_tracking_xlsx(demo_con):
    name, data = reports.shipment_tracking_xlsx(demo_con, TW_KEY)
    assert name == "Matalan Carton Shipment Tracking_NPS_BD Aug-26 to Aug-26 (Transworld).xlsx"
    wb = openpyxl.load_workbook(BytesIO(data))
    assert wb.sheetnames == ["TRACK_STD", "TRACK_CUS"]
    assert wb["TRACK_STD"].max_row == 56 and wb["TRACK_CUS"].max_row == 92
    assert wb["TRACK_STD"]["A1"].value == "PO Received Date"
    assert reports.shipment_tracking_xlsx(demo_con, "Nobody") is None


def pdf_text(data):
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(BytesIO(data)).pages)


def test_full_tracker_pdf_pacd(demo_con):
    name, data = reports.full_tracker_pdf(demo_con, TW_KEY, 2026, [3], "PacD")
    assert name == "Full Tracker Sheet for PacD_Transworld Packaging (BD) Ltd_Q3_2026.pdf"
    text = pdf_text(data)
    total = sum(r["total_amount"] for kind in ("STD", "CUS") for r in core.rows(demo_con, kind, supplier=TW_KEY))
    for expected in ("Full Tracker Report (Transworld Packaging (BD) Ltd.)", "Standard Carton",
                     "Q3 2026 (Jul-Aug) Quarter", "Q3 2026 (Jul-Aug) Subtotal", "Standard Carton Total",
                     "Custom Carton Total", "Tracker Grand Total", f"{total:,.2f}", "PACD",
                     "Page 1 of", "Generated:"):
        assert expected in text, expected
    assert len(pypdf.PdfReader(BytesIO(data)).pages[0].images) == 2  # both logos


def test_full_tracker_pdf_matalan_has_no_pacd(demo_con):
    name, data = reports.full_tracker_pdf(demo_con, UG_KEY, 2026, [1, 2, 3], "Matalan")
    assert name == "Full Tracker Sheet for Matalan_Uniglory Packaging Industries Ltd_Q1-Q3_2026.pdf"
    text = pdf_text(data)
    assert "PACD" not in text
    assert "Rebates" in text and "Tracker Grand Total" in text


def test_full_tracker_pdf_without_rows_is_none(demo_con):
    assert reports.full_tracker_pdf(demo_con, TW_KEY, 2026, [1], "PacD") is None


SUMMARY = "Matalan Carton Shipment Tracking_Summary_Aug-Aug 26.xlsx"
TW_FILES = [
    "Transworld Packaging (BD) Ltd/Full Tracker Sheet for PacD_Transworld Packaging (BD) Ltd_Q3_2026.pdf",
    "Transworld Packaging (BD) Ltd/Full Tracker Sheet for Matalan_Transworld Packaging (BD) Ltd_Q3_2026.pdf",
    "Transworld Packaging (BD) Ltd/Tracker Raw Data_Transworld Packaging (BD) Ltd_Q3_2026.xlsx",
    "Transworld Packaging (BD) Ltd/Matalan Carton Shipment Tracking_NPS_BD Aug-26 to Aug-26 (Transworld).xlsx",
]
UG_FILES = [
    "Uniglory Packaging Industries Ltd/Full Tracker Sheet for PacD_Uniglory Packaging Industries Ltd_Q3_2026.pdf",
    "Uniglory Packaging Industries Ltd/Full Tracker Sheet for Matalan_Uniglory Packaging Industries Ltd_Q3_2026.pdf",
    "Uniglory Packaging Industries Ltd/Tracker Raw Data_Uniglory Packaging Industries Ltd_Q3_2026.xlsx",
    "Uniglory Packaging Industries Ltd/Matalan Carton Shipment Tracking_NPS_BD Aug-26 to Aug-26 (Uniglory).xlsx",
]


def test_save_all(demo_con, tmp_path):
    root = tmp_path / "Reports"
    paths = reports.save_all(demo_con, [TW_KEY], 2026, [3], root)
    assert sorted(p.relative_to(root).as_posix() for p in paths) == sorted(
        f"August/{p}" for p in [SUMMARY] + TW_FILES)
    assert all(p.stat().st_size > 0 for p in paths)
    assert reports.save_all(demo_con, [TW_KEY], 2026, [1], root) == []


def test_build_and_save_several_suppliers(demo_con, tmp_path):
    items = reports.build_all(demo_con, [TW_KEY, UG_KEY], 2026, [3])
    assert [path for path, _ in items] == [SUMMARY] + TW_FILES + UG_FILES
    with zipfile.ZipFile(BytesIO(reports.to_zip(items))) as z:
        assert z.namelist() == [SUMMARY] + TW_FILES + UG_FILES
        assert z.read(SUMMARY) == items[0][1]
    root = tmp_path / "Reports"
    paths = reports.save_all(demo_con, [TW_KEY, UG_KEY], 2026, [3], root)
    assert sorted(p.relative_to(root).as_posix() for p in paths) == sorted(
        f"August/{p}" for p in [SUMMARY] + TW_FILES + UG_FILES)
    assert reports.month_name(demo_con, [TW_KEY, UG_KEY], 2026, [3]) == "August"


def test_supplier_without_rows_in_selection_is_skipped(demo_con):
    insert(demo_con, "STD", shipment_date="2025-01-15", packaging_supplier="Old Supplier Ltd.",
           supplier_key="OLD", carton_qty=1, unit_price=1.0, total_amount=1.0)
    paths = [path for path, _ in reports.build_all(demo_con, [TW_KEY, "OLD"], 2026, [3])]
    assert paths == [SUMMARY] + TW_FILES


def test_app_smoke(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    core.save(core.connect(db), core.read_workbook(DEMO), DEMO.name)
    at = AppTest.from_file(str(HERE / "pages" / "1_Dashboard.py"), default_timeout=30).run()
    assert not at.exception
    assert [item.value for item in at.subheader] == [
        "Quarterly carton sales by packaging supplier",
        "Standard and custom cartons",
        "Live Summary sheet",
        "Tracker and eBiz cross-reference",
    ]
    assert at.selectbox[0].label == "Fiscal year"
    assert at.multiselect[0].label == "Packaging suppliers"


def test_app_builds_reports_for_every_selected_supplier(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    core.save(core.connect(db), core.read_workbook(DEMO), DEMO.name)
    at = AppTest.from_file(str(HERE / "pages" / "2_Reports.py"), default_timeout=60).run()
    assert at.multiselect[0].options == [TW, UG]
    assert at.multiselect[0].value == [TW_KEY, UG_KEY]
    next(b for b in at.button if b.label == "Build reports").click().run()
    assert not at.exception
    _, items, zipped, zip_name = at.session_state["built_reports"]
    assert [path for path, _ in items] == [SUMMARY] + TW_FILES + UG_FILES
    assert zip_name == "Matalan Carton Reports_August 2026.zip"
    assert zipfile.ZipFile(BytesIO(zipped)).namelist() == [SUMMARY] + TW_FILES + UG_FILES


class FakeUpload:
    """Stands in for a Streamlit UploadedFile (AppTest can't drive the file uploader)."""

    def __init__(self, path):
        self.name, self._data = path.name, path.read_bytes()

    def getvalue(self):
        return self._data


def test_app_date_check_imports_and_remembers_suggested_dates(tmp_path, monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    path = make_xlsx(tmp_path / "a.xlsx", std=[rec(ship="2025-12-21", po_received_date=datetime(2025, 6, 12)), rec()])
    monkeypatch.setattr(st, "file_uploader", lambda *a, **k: [] if k.get("key") == "ebiz_upload" else [FakeUpload(path)])
    at = AppTest.from_file(str(HERE / "pages" / "5_Import.py"), default_timeout=30).run()
    assert not at.exception
    assert at.warning[0].value.startswith("1 row(s) have dates that need review.")
    next(b for b in at.button if b.label == "Confirm tracker import").click().run()
    assert not at.exception
    con = core.connect(db)
    assert [r[0] for r in con.execute("SELECT po_received_date FROM track_std ORDER BY id")] == ["2025-12-06", None]
    at.run()  # same file still uploaded: the saved correction is used, nothing to review
    assert "1 row(s) used a previously saved date correction." in [i.value for i in at.info]
    assert not at.warning
    assert len(core.saved_dates(con)) == 1


def test_app_no_data_for_selection(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    con = core.connect(db)
    core.save(con, core.read_workbook(DEMO), DEMO.name)
    # Insert a row for Old Supplier with Q1 2025 data (shipment_date 2025-01-15)
    insert(con, "STD", shipment_date="2025-01-15", packaging_supplier="Old Supplier Ltd.", supplier_key="OLD",
           carton_qty=100, unit_price=1.0, total_amount=100.0)
    at = AppTest.from_file(str(HERE / "pages" / "2_Reports.py"), default_timeout=30).run()
    # Select only the old supplier, which has no 2026 data
    at.multiselect[0].set_value(["OLD"]).run()
    at.selectbox[0].select(2026).run()
    # Verify "No data for this selection." appears
    assert "No data for this selection." in [e.value for e in at.info]


def make_ebiz(path, rows, header_row=1, drop=None):
    """Write an eBiz export: headers in header_row, rows (dicts keyed by core.EBIZ_HEADERS names) below."""
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = [c for c in core.EBIZ_HEADERS if c != drop]
    for i, c in enumerate(cols, 1):
        ws.cell(header_row, i, core.EBIZ_HEADERS[c])
    for r, row in enumerate(rows, header_row + 1):
        for i, c in enumerate(cols, 1):
            ws.cell(r, i, row.get(c))
    wb.save(path)
    return path


def erow(order="7001", qty=100, rebate=1.0, vendor="Pack A Ltd.", **kw):
    """One eBiz export row for make_ebiz."""
    return {"vendor": vendor, "vendor_code": "M0001", "product": "MLN-CARTONA1", "order_number": order,
            "qty_shipped": qty, "net_price_1000": 840, "gross_price_1000": 880, "total_rebate": rebate,
            "currency": "USD", **kw}


def test_read_ebiz_splits_order_numbers(tmp_path):
    path = make_ebiz(tmp_path / "e.xlsx", [erow("7001|7002"), erow(7003.0, qty=5), erow(None, qty=None),
                                           erow(" 7004 | ")], header_row=2)
    rows = core.read_ebiz(path)
    assert [(r["order_number"], r["so_1"], r["so_2"]) for r in rows] == [
        ("7001|7002", "7001", "7002"), ("7003", "7003", None), ("7004 |", "7004", None)]
    assert rows[0]["vendor_key"] == "PACK"
    assert rows[1]["qty_shipped"] == 5.0 and rows[1]["currency"] == "USD"


def test_read_ebiz_skips_vendorless_total_lines(tmp_path):
    total = erow(None, qty=10, vendor=None)  # a grand-total line: no Vendor, no Order Number
    files = {"a.xlsx": core.read_ebiz(make_ebiz(tmp_path / "a.xlsx", [erow("1"), total])),
             "b.xlsx": core.read_ebiz(make_ebiz(tmp_path / "b.xlsx", [erow("2", vendor="Other Co"), total]))}
    assert [r["vendor_key"] for rows in files.values() for r in rows] == ["PACK", "OTHER"]
    assert core.ebiz_clashes(files) == []
    con = core.connect(tmp_path / "t.db")
    assert core.replace_ebiz(con, files) == {"a.xlsx": 1, "b.xlsx": 1}


def test_read_ebiz_needs_its_headers(tmp_path):
    with pytest.raises(core.ImportFileError, match="header row not found"):
        core.read_ebiz(make_ebiz(tmp_path / "a.xlsx", [erow()], drop="order_number"))
    with pytest.raises(core.ImportFileError, match="Qty Shipped"):
        core.read_ebiz(make_ebiz(tmp_path / "b.xlsx", [erow()], drop="qty_shipped"))
    assert core.read_ebiz(make_ebiz(tmp_path / "c.xlsx", [erow()], drop="currency"))[0]["currency"] is None


def test_ebiz_preview(tmp_path):
    rows = core.read_ebiz(make_ebiz(tmp_path / "a.xlsx", [erow("1", qty=10, rebate=1.5),
                                                          erow("2", qty=5, rebate=0.25, vendor="PACK A LIMITED")]))
    assert core.ebiz_preview({"a.xlsx": rows}) == [{
        "File": "a.xlsx", "Vendor": "PACK", "Spellings in file": "PACK A LIMITED / Pack A Ltd.", "Rows": 2,
        "Qty Shipped": 15.0, "Total Rebate": 1.75}]


def test_ebiz_upload_replaces_all_ebiz_data(tmp_path):
    con = core.connect(tmp_path / "t.db")
    assert core.ebiz_status(con) is None
    a = {"a.xlsx": core.read_ebiz(make_ebiz(tmp_path / "a.xlsx", [erow("1"), erow("2", vendor="Other Co")]))}
    assert core.replace_ebiz(con, a) == {"a.xlsx": 2}
    assert core.ebiz_vendors(con) == ["OTHER", "PACK"]
    b = {"b.xlsx": core.read_ebiz(make_ebiz(tmp_path / "b.xlsx", [erow("3")]))}
    core.replace_ebiz(con, b)
    assert [r[0] for r in con.execute("SELECT order_number FROM ebiz")] == ["3"]
    count, files, imported_at = core.ebiz_status(con)
    assert (count, files) == (1, ["b.xlsx"]) and imported_at
    clash = {"c.xlsx": b["b.xlsx"], "d.xlsx": b["b.xlsx"]}
    assert core.ebiz_clashes(clash) == [{"Vendor": "PACK", "Files": "c.xlsx / d.xlsx"}]
    with pytest.raises(ValueError, match="PACK"):
        core.replace_ebiz(con, clash)
    assert core.ebiz_status(con)[0] == 1  # a refused upload keeps the old data


def test_old_database_gets_the_ebiz_table(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-08-03", carton_qty=1, unit_price=1.0, total_amount=1.0)
    con.execute("DROP TABLE ebiz")
    con.execute("PRAGMA user_version = 2")
    con.close()
    con = core.connect(tmp_path / "t.db")
    assert core.ebiz_status(con) is None
    assert con.execute("SELECT COUNT(*) FROM track_std").fetchone()[0] == 1


@pytest.fixture
def xref_con(tmp_path):
    """Supplier PACK: tracker SOs 1001-1004 and rows without an SO#; eBiz rows giving every status."""
    con = core.connect(tmp_path / "t.db")
    for so, qty, ship in [("1001", 100, "2026-08-03"), ("1002", 100, "2026-08-04"), ("1003", 50, "2026-08-05"),
                          ("1004", 10, "2026-08-06"), ("1004", 1, None), (None, 5, "2026-08-07"),
                          ("0", 2, "2026-08-08")]:
        insert(con, "STD", shipment_date=ship, packaging_supplier="Pack A Ltd.", supplier_key="PACK", ebiz_so=so,
               carton_qty=qty, unit_price=1.0, total_amount=float(qty))
    rows = core.read_ebiz(make_ebiz(tmp_path / "e.xlsx", [
        erow("1001", qty=100, rebate=8.8), erow("9999|1002", qty=80, rebate=4.0), erow("1003", qty=70, rebate=2.0),
        erow("8888|8889", qty=7, rebate=0.5), erow("1001", qty=3, vendor="Other Co")]))
    core.replace_ebiz(con, {"e.xlsx": rows})
    return con


def test_cross_reference(xref_con):
    x = core.cross_reference(xref_con, "PACK")
    got = [(o["so"], o["tracker_qty"], round(o["tracker_rebate"], 4), o["ebiz_qty"], o["ebiz_rebate"],
            o["status"], o["via_so2"]) for o in x["orders"]]
    assert got == [("1001", 100, 8.8, 100, 8.8, "Match", False),  # rebate 5 + PACD (100 - 5) * 4%
                   ("1002", 100, 8.8, 80, 4.0, "Tracker > eBiz", True),  # matched through SO_2
                   ("1003", 50, 4.4, 70, 2.0, "Tracker < eBiz", False),
                   ("1004", 11, 0.88, 0, 0.0, "Not in eBiz", False)]  # undated row counts, rebate 0
    assert [r["so"] for r in x["rows"]["STD"]] == ["1001", "1002", "1003", "1004", "1004", None, None]
    assert x["rows"]["CUS"] == []
    assert [(e["key"], e["qty_shipped"]) for e in x["ebiz_only"]] == [("8888", 7.0)]  # Other Co ignored
    assert core.cross_reference(xref_con, "NOBODY") == {"orders": [], "rows": {"STD": [], "CUS": []},
                                                        "ebiz_only": []}
    assert sorted(["12", None, "AB", "3", "²"], key=core.so_order) == ["3", "12", "AB", "²", None]  # "²".isdigit()


def test_cross_reference_overview_drives_live_dashboard(xref_con):
    rows = core.cross_reference_overview(xref_con)
    pack = next(row for row in rows if row["supplier_key"] == "PACK")
    assert pack == {
        "supplier_key": "PACK",
        "packaging_supplier": "Pack A Ltd.",
        "tracker_orders": 4,
        "matches": 1,
        "not_in_ebiz": 1,
        "quantity_mismatches": 2,
        "ebiz_only": 1,
    }


DISCREPANCY_TITLES = ("Tracker Orders Not Found in eBiz", "Quantity Mismatch (Tracker > eBiz)",
                      "Quantity Mismatch (Tracker < eBiz)", "Tracker Rows With No SO#",
                      "eBiz Orders Not Found in Tracker")


def test_discrepancy_pdf(xref_con):
    name, data = reports.discrepancy_pdf(xref_con, "PACK", core.cross_reference(xref_con, "PACK"))
    assert name == f"Discrepancy Report_Pack A Ltd_{date.today():%d-%b-%Y}.pdf"
    text = pdf_text(data)
    for expected in ("Discrepancy Report (Pack A Ltd.)", *DISCREPANCY_TITLES, "Standard Carton", "eBiz QTY",
                     "8888|8889", "Page 1 of", "Generated:"):
        assert expected in text, expected
    assert "No discrepancies found." not in text
    assert len(pypdf.PdfReader(BytesIO(data)).pages[0].images) == 2  # both logos


def test_discrepancy_pdf_says_when_a_section_is_empty(tmp_path):
    con = core.connect(tmp_path / "t.db")
    insert(con, "STD", shipment_date="2026-08-03", packaging_supplier="Pack A Ltd.", supplier_key="PACK",
           ebiz_so="1001", carton_qty=100, unit_price=1.0, total_amount=100.0)
    core.replace_ebiz(con, {"e.xlsx": core.read_ebiz(make_ebiz(tmp_path / "e.xlsx", [erow("1001", qty=100)]))})
    _, data = reports.discrepancy_pdf(con, "PACK", core.cross_reference(con, "PACK"))
    text = pdf_text(data)
    assert all(t in text for t in DISCREPANCY_TITLES)
    assert text.count("No discrepancies found.") == 5


def test_discrepancy_pdf_section_bar_never_ends_a_page(tmp_path):
    con = core.connect(tmp_path / "t.db")
    for so, qty in [(str(1000 + i), 1) for i in range(42)] + [("9000", 10)]:  # 42 not in eBiz, 1 mismatch
        insert(con, "STD", shipment_date="2026-08-03", packaging_supplier="Pack A Ltd.", supplier_key="PACK",
               ebiz_so=so, carton_qty=qty, unit_price=1.0, total_amount=float(qty))
    rows = core.read_ebiz(make_ebiz(tmp_path / "e.xlsx", [erow("9000", qty=5), erow("8888", qty=7)]))
    core.replace_ebiz(con, {"e.xlsx": rows})  # the mismatch and an eBiz-only order give two sections a table
    full = core.cross_reference(con, "PACK")
    for n in range(23, 42):  # the first table ends near the page foot (about 24 rows fit), so bars land there
        xref = {**full, "orders": full["orders"][:n] + full["orders"][-1:],
                "rows": {"STD": full["rows"]["STD"][:n] + full["rows"]["STD"][-1:], "CUS": []}}
        _, data = reports.discrepancy_pdf(con, "PACK", xref)
        for page in pypdf.PdfReader(BytesIO(data)).pages:
            above_footer = page.extract_text().splitlines()[-2]  # the last line is "Page x of y"
            assert above_footer not in DISCREPANCY_TITLES, (n, above_footer)


XREF_HEADER = ["SO#", "Tracker QTY", "Tracker Rebate", "eBiz QTY", "eBiz Rebate", "QTY Difference", "Status"]


def test_cross_reference_xlsx(xref_con):
    name, data = reports.cross_reference_xlsx(xref_con, "PACK", core.cross_reference(xref_con, "PACK"))
    assert name == f"Cross Reference_Pack A Ltd_{date.today():%d-%b-%Y}.xlsx"
    ws, rows = summary_sheet(data)
    assert rows[0][0] == "Cross Reference"
    assert len(ws._images) == 2  # PacD and Matalan logos
    head = rows.index(XREF_HEADER)
    assert rows[head + 1:head + 7] == [
        [1001, 100, 8.8, 100, 8.8, 0, "Match"],  # all-digit SO#s are numbers: no "stored as text" triangles
        [1002, 100, 8.8, 80, 4.0, 20, "Tracker > eBiz"],
        [1003, 50, 4.4, 70, 2.0, -20, "Tracker < eBiz"],
        [1004, 11, 0.88, 0, 0, 11, "Not in eBiz"],
        ["(No SO#)", 7, 0.62, None, None, 7, None],  # QTY Difference adds up to the Total line
        ["Total", 268, 23.5, 250, 14.8, 18, None]]
    assert ws.page_setup.orientation == "portrait" and ws.print_title_rows == f"${head + 1}:${head + 1}"
    assert ws.oddFooter.center.text == "Page &P of &N"
    # The eBiz-only list has its own sheet: a sheet's print titles repeat on all its pages, so on this sheet they
    # would put the tracker header above the list.
    wb = openpyxl.load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Cross Reference", "eBiz not in tracker"] and not ws.row_breaks.brk
    assert ["Order", "eBiz QTY", "eBiz Rebate", None, None, None, None] not in rows
    only_ws = wb["eBiz not in tracker"]
    only = [[c.value for c in r] for r in only_ws.iter_rows()]
    first = only.index(["Order", "eBiz QTY", "eBiz Rebate", None, None, None, None])
    assert only[first + 1:first + 3] == [[8888, 7, 0.5, None, None, None, None],
                                         ["Total", 7, 0.5, None, None, None, None]]
    assert only_ws.print_title_rows == f"${first + 1}:${first + 1}" and len(only_ws._images) == 2


def test_build_all_adds_discrepancy_files_for_suppliers_with_ebiz(demo_con, tmp_path):
    so = demo_con.execute("SELECT ebiz_so FROM track_std WHERE supplier_key = ? AND ebiz_so IS NOT NULL",
                          (TW_KEY,)).fetchone()[0]
    core.replace_ebiz(demo_con, {"e.xlsx": core.read_ebiz(make_ebiz(tmp_path / "e.xlsx", [erow(so, vendor=TW)]))})
    folder = "Transworld Packaging (BD) Ltd"
    today = f"{date.today():%d-%b-%Y}"
    disc = [f"{folder}/Discrepancy Report_{folder}_{today}.pdf", f"{folder}/Cross Reference_{folder}_{today}.xlsx"]
    items = reports.build_all(demo_con, [TW_KEY, UG_KEY], 2026, [3], discrepancy=True)
    assert [p for p, _ in items] == [SUMMARY] + TW_FILES + disc + UG_FILES  # Uniglory has no eBiz rows
    assert [p for p, _ in reports.build_all(demo_con, [TW_KEY, UG_KEY], 2026, [3])] == [SUMMARY] + TW_FILES + UG_FILES
    # all tracker rows are compared, so a supplier with nothing in the chosen quarters still gets them
    assert [p for p, _ in reports.build_all(demo_con, [TW_KEY], 2026, [1], discrepancy=True)] == [SUMMARY] + disc
    root = tmp_path / "Reports"
    paths = reports.save_all(demo_con, [TW_KEY], 2026, [3], root, discrepancy=True)
    assert sorted(p.relative_to(root).as_posix() for p in paths) == sorted(
        f"August/{p}" for p in [SUMMARY] + TW_FILES + disc)


def test_app_replaces_ebiz_data(tmp_path, monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    core.save(core.connect(db), core.read_workbook(DEMO), DEMO.name)
    path = make_ebiz(tmp_path / "ebiz.xlsx", [erow("1", vendor=TW), erow("2", vendor="Nobody Ltd")])
    monkeypatch.setattr(st, "file_uploader",
                        lambda *a, **k: [FakeUpload(path)] if k.get("key") == "ebiz_upload" else [])
    at = AppTest.from_file(str(HERE / "pages" / "5_Import.py"), default_timeout=30).run()
    assert not at.exception
    assert any("NOBODY" in w.value for w in at.warning)  # vendor without tracker rows
    next(b for b in at.button if b.label == "Replace all eBiz data").click().run()
    assert not at.exception
    assert core.ebiz_status(core.connect(db))[0] == 2
    report = AppTest.from_file(str(HERE / "pages" / "2_Reports.py"), default_timeout=30).run()
    checkbox = next(c for c in report.checkbox if c.label == "Discrepancy report and cross reference")
    assert checkbox.value and not checkbox.disabled
    assert not any("current eBiz data will be removed" in w.value for w in at.warning)


def test_app_warns_which_vendors_a_replace_drops(tmp_path, monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    db = tmp_path / "app.db"
    monkeypatch.setenv("CARTON_DB", str(db))
    con = core.connect(db)
    core.save(con, core.read_workbook(DEMO), DEMO.name)
    old = {name: core.read_ebiz(make_ebiz(tmp_path / name, [erow("1", vendor=v)]))
           for name, v in [("tw.xlsx", TW), ("other.xlsx", "Other Co"), ("nobody.xlsx", "Nobody Ltd")]}
    core.replace_ebiz(con, old)
    path = make_ebiz(tmp_path / "new.xlsx", [erow("2", vendor=TW)])
    monkeypatch.setattr(st, "file_uploader",
                        lambda *a, **k: [FakeUpload(path)] if k.get("key") == "ebiz_upload" else [])
    at = AppTest.from_file(str(HERE / "pages" / "5_Import.py"), default_timeout=30).run()
    assert not at.exception
    assert "Current eBiz data will be removed for: NOBODY, OTHER." in [w.value for w in at.warning]
