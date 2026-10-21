"""SQLite storage, Excel import and report queries for the Matalan carton tracker."""
import hashlib
import json
import math
import os
import re
import sqlite3
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

import openpyxl

DEFAULT_DB = Path(__file__).with_name("carton.db")
SCHEMA_VERSION = 4
TABLES = {"STD": "track_std", "CUS": "track_cus"}
SHEETS = {"STD": "TRACK_STD", "CUS": "TRACK_CUS"}

# Database column -> header text in the tracker sheets, in sheet order.
HEADERS = {
    "po_received_date": "PO Received Date",
    "shipment_date": "Shipment Date",
    "matalan_po": "Matalan Product PO#",
    "supplier_name": "Supplier Name",
    "factory_name": "Factory Name",
    "packaging_supplier": "Packaging Supplier Name",
    "ebiz_customer_ref": "eBiz Customer Reference #",
    "ebiz_so": "eBiz Sales Order #",
    "carton_qty": "Carton QTY",
    "carton_style": "Carton Style",
    "flute": "FLUTE",
    "carton_selection": "Carton Selection",
    "length_mm": "Length(mm)",
    "width_mm": "Width(mm)",
    "height_mm": "Height(mm)",
    "l_mm": "L",
    "w_mm": "W",
    "area_sqm": "Area(sqm)",
    "unit_price": "Carton's Unit Price (CNY,EXW w/VAT)",
    "final_price": "Final Price, CNY/CTN",
    "total_amount": "Total Amount",
    "note": "NOTE",
}
COLUMNS = list(HEADERS)
DATE_COLS = {"po_received_date", "shipment_date"}
INT_COLS = {"carton_qty"}
REAL_COLS = {"length_mm", "width_mm", "height_mm", "l_mm", "w_mm", "area_sqm",
             "unit_price", "final_price", "total_amount"}
# Columns that identify a tracker row across imports (prices are formulas, so left out).
KEY_COLS = [c for c in COLUMNS if c not in {"area_sqm", "unit_price", "final_price", "total_amount"}]

FIRST_DATE = date(2020, 1, 1)  # earlier dates are typos
LEAD_DAYS = 60  # a day/month swap is suggested only if it puts the PO at most this many days before shipment

# eBiz export column -> header text (first sheet of each vendor's export).
EBIZ_HEADERS = {
    "vendor": "Vendor",
    "vendor_code": "Vendor Code",
    "product": "Product",
    "order_number": "Order Number",
    "qty_shipped": "Qty Shipped",
    "net_price_1000": "Current Net Price per 1000",
    "gross_price_1000": "Current Gross Price per 1000",
    "total_rebate": "Total Rebate",
    "currency": "Currency",
}
EBIZ_REQUIRED = ["vendor", "order_number", "qty_shipped", "total_rebate"]
EBIZ_REAL = {"qty_shipped", "net_price_1000", "gross_price_1000", "total_rebate"}
EBIZ_COLUMNS = list(EBIZ_HEADERS) + ["so_1", "so_2", "vendor_key"]

SEED_RATES = [
    ("STD", "2025-12-01", 0.04, 0.05),
    ("CUS", "2025-12-01", 0.075, 0.10),
    ("STD", "2026-07-01", 0.04, 0.05),
    ("CUS", "2026-07-01", 0.075, 0.05),
]

VIEW_SQL = """
CREATE VIEW {view} AS
SELECT t.*, r.pacd_pct, r.matalan_pct,
  t.unit_price * t.carton_qty * r.matalan_pct AS matalan_rebate,
  (t.total_amount - t.unit_price * t.carton_qty * r.matalan_pct) * r.pacd_pct AS pacd,
  CAST(strftime('%Y', t.shipment_date) AS INTEGER) + (strftime('%m', t.shipment_date) = '12') AS fiscal_year,
  CASE
    WHEN t.shipment_date IS NULL THEN NULL
    WHEN CAST(strftime('%m', t.shipment_date) AS INTEGER) IN (12, 1, 2, 3) THEN 1
    WHEN CAST(strftime('%m', t.shipment_date) AS INTEGER) BETWEEN 4 AND 6 THEN 2
    WHEN CAST(strftime('%m', t.shipment_date) AS INTEGER) IN (7, 8) THEN 3
    ELSE 4
  END AS quarter
FROM {table} t
LEFT JOIN rates r ON r.carton_type = '{kind}' AND r.effective_from = (
  SELECT MAX(effective_from) FROM rates
  WHERE carton_type = '{kind}' AND effective_from <= t.shipment_date);
"""


def _col_type(col):
    return "INTEGER" if col in INT_COLS else "REAL" if col in REAL_COLS else "TEXT"


def _schema():
    cols = ",\n  ".join(f"{c} {_col_type(c)}" for c in COLUMNS)
    sql = [f"CREATE TABLE IF NOT EXISTS {t} (\n  id INTEGER PRIMARY KEY,\n  {cols},\n"
           f"  supplier_key TEXT,\n  source_file TEXT,\n  imported_at TEXT\n);" for t in TABLES.values()]
    sql.append("""CREATE TABLE IF NOT EXISTS rates (
  carton_type TEXT NOT NULL CHECK (carton_type IN ('STD', 'CUS')),
  effective_from TEXT NOT NULL,
  pacd_pct REAL NOT NULL,
  matalan_pct REAL NOT NULL,
  PRIMARY KEY (carton_type, effective_from)
);""")
    sql.append("""CREATE TABLE IF NOT EXISTS date_fixes (
  row_key TEXT PRIMARY KEY,
  carton_type TEXT NOT NULL,
  packaging_supplier TEXT,
  ebiz_so TEXT,
  carton_qty INTEGER,
  file_po TEXT,
  file_ship TEXT,
  po_received_date TEXT,
  shipment_date TEXT,
  saved_at TEXT NOT NULL
);""")
    ebiz_cols = ",\n  ".join(f"{c} {'REAL' if c in EBIZ_REAL else 'TEXT'}" for c in EBIZ_COLUMNS)
    sql.append(f"CREATE TABLE IF NOT EXISTS ebiz (\n  id INTEGER PRIMARY KEY,\n  {ebiz_cols},\n"
               f"  source_file TEXT,\n  imported_at TEXT\n);")
    sql.append("""CREATE TABLE IF NOT EXISTS app_meta (
  key TEXT PRIMARY KEY,
  value TEXT
);""")
    sql.append("""CREATE TABLE IF NOT EXISTS packaging_suppliers (
  id INTEGER PRIMARY KEY,
  name TEXT NOT NULL COLLATE NOCASE UNIQUE,
  created_at TEXT NOT NULL
);""")
    sql.append("""CREATE TABLE IF NOT EXISTS assignment_factories (
  id INTEGER PRIMARY KEY,
  name TEXT NOT NULL COLLATE NOCASE UNIQUE,
  created_at TEXT NOT NULL
);""")
    sql.append("""CREATE TABLE IF NOT EXISTS factory_assignments (
  factory_id INTEGER NOT NULL REFERENCES assignment_factories(id) ON DELETE CASCADE,
  packaging_supplier_id INTEGER NOT NULL REFERENCES packaging_suppliers(id) ON DELETE RESTRICT,
  priority INTEGER NOT NULL,
  PRIMARY KEY (factory_id, packaging_supplier_id),
  UNIQUE (factory_id, priority)
);""")
    for kind, table in TABLES.items():
        view = f"v_{kind.lower()}"
        sql.append(f"DROP VIEW IF EXISTS {view};")
        sql.append(VIEW_SQL.format(view=view, table=table, kind=kind))
    return "\n".join(sql)


def connect(path=None):
    """Open the database; create the schema when its version is not SCHEMA_VERSION.

    Path: the argument, else env CARTON_DB, else carton.db next to this file.
    """
    con = sqlite3.connect(path or os.environ.get("CARTON_DB") or DEFAULT_DB)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    if con.execute("PRAGMA user_version").fetchone()[0] != SCHEMA_VERSION:
        con.executescript(_schema())
        con.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")
    if not con.execute("SELECT 1 FROM rates LIMIT 1").fetchone():
        with con:
            con.executemany("INSERT INTO rates VALUES (?, ?, ?, ?)", SEED_RATES)
    sync_assignment_catalog(con)
    return con


EXCEL_ERRORS = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


class ImportFileError(ValueError):
    """The workbook can't be imported. The message says why."""


def _norm(header):
    return re.sub(r"\s+", " ", str(header or "")).strip().lower()


def _clean(value):
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value in EXCEL_ERRORS:
            return None
    return value


def supplier_key(name):
    """Upper-cased first word of a packaging supplier name, so its spellings merge (port of the VBA FirstWord).

    'M&U  Packaging Ltd.' -> 'M&U'. None when there is no word.
    """
    if name is None:
        return None
    for ch in '.,;:"\'!?':
        name = name.replace(ch, " ")
    words = name.split()
    return words[0].upper() if words else None


_TEXT_DATE = re.compile(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})")


def _text_date(text):
    """'18/12/2025' -> date(2025, 12, 18). Text dates are read day first, as the factories type them."""
    m = _TEXT_DATE.fullmatch(text.strip())
    if not m:
        return None
    day, month, year = map(int, m.groups())
    try:
        return date(year + 2000 if year < 100 else year, month, day)
    except ValueError:
        return None


def _convert(col, value):
    if value is None:
        return None
    if col in DATE_COLS:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        d = _text_date(value) if isinstance(value, str) else None
        return d.isoformat() if d else None
    if col in INT_COLS or col in REAL_COLS:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        return int(round(number)) if col in INT_COLS else number
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _shown(value):
    """A date cell as the user sees it: a date as 07-Dec-2025 (no day/month doubt), text as typed."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return f"{value:%d-%b-%Y}"
    return None if value is None else str(value)


def _row_key(kind, raw, row):
    """Hash of the row's input cells, so the same tracker row has the same key in next month's file."""
    ident = [kind] + [_shown(raw[c]) if c in DATE_COLS else row[c] for c in KEY_COLS]
    return hashlib.sha1(json.dumps(ident).encode()).hexdigest()


def _find_columns(sheet, ws):
    """Return (column index per COLUMNS, header row number).

    The header row is the first of rows 1-10 with a 'Shipment Date' cell: row 2 in the
    master tracker (title above), row 1 in the per-supplier exports.
    """
    for n, header_row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        found = {_norm(h): i for i, h in enumerate(header_row)}
        if _norm(HEADERS["shipment_date"]) in found:
            break
    else:
        raise ImportFileError(f"{sheet}: header row not found (no 'Shipment Date' cell in rows 1-10).")
    index = {}
    for col in COLUMNS:
        index[col] = found.get(_norm(HEADERS[col]))
        if index[col] is None and col == "unit_price":
            index[col] = next((i for h, i in found.items() if h.startswith("carton's unit price")), None)
        if index[col] is None:
            raise ImportFileError(f"{sheet}: column '{HEADERS[col]}' not found in the header row (row {n}).")
    return index, n


def read_workbook(src):
    """Read the data rows of TRACK_STD and TRACK_CUS from a path or binary file object.

    Returns {"STD": [...], "CUS": [...]}. Each row is a dict of COLUMNS plus "supplier_key",
    "_row" (Excel row number), "_file" (the two date cells as they are in the file) and
    "_key" (identifies the row in later imports, for saved date corrections).
    """
    try:
        wb = openpyxl.load_workbook(src, data_only=True)
    except Exception as e:  # not an xlsx, corrupt or password protected
        raise ImportFileError(f"Can't open the workbook: {e}") from e
    names = {n.upper(): n for n in wb.sheetnames}
    for sheet in SHEETS.values():
        if sheet not in names:
            raise ImportFileError(f"Sheet '{sheet}' not found.")
    data = {}
    for kind, sheet in SHEETS.items():
        ws = wb[names[sheet]]
        index, header = _find_columns(sheet, ws)
        rows = []
        for n, values in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), start=header + 1):
            raw = {c: _clean(values[i]) if i < len(values) else None for c, i in index.items()}
            if raw["shipment_date"] is None and raw["carton_qty"] is None:
                continue  # template row: formulas but no input
            row = {c: _convert(c, v) for c, v in raw.items()}
            row["supplier_key"] = supplier_key(row["packaging_supplier"])
            row["_row"] = n
            row["_file"] = {c: raw[c] for c in DATE_COLS}
            row["_key"] = _row_key(kind, raw, row)
            rows.append(row)
        if rows and all(r["total_amount"] is None for r in rows):
            raise ImportFileError(f"{sheet}: the file has no calculated values. "
                                  "Open it in Excel, save, and import again.")
        data[kind] = rows
    return data


def _month(row):
    return row["shipment_date"][:7] if row["shipment_date"] else None


_SAME_KEY = "supplier_key IS ? AND substr(shipment_date, 1, 7) IS ?"


def preview(con, data):
    """Rows per sheet, supplier key and shipment month, and how many DB rows each replaces."""
    out = []
    for kind, rows in data.items():
        groups = {}
        for r in rows:
            groups.setdefault((r["supplier_key"], _month(r)), []).append(r["packaging_supplier"])
        for (key, month), names in sorted(groups.items(), key=lambda g: (str(g[0][0]), str(g[0][1]))):
            replaced = con.execute(f"SELECT COUNT(*) FROM {TABLES[kind]} WHERE {_SAME_KEY}",
                                   (key, month)).fetchone()[0]
            out.append({"Sheet": SHEETS[kind], "Supplier": key,
                        "Spellings in file": " / ".join(sorted(set(names) - {None})),
                        "Shipment month": month or "no valid date", "Rows in file": len(names),
                        "Rows replaced in DB": replaced})
    return out


def _swap(d):
    """d with day and month exchanged, or None if that is not a date."""
    try:
        return d.replace(month=d.day, day=d.month)
    except ValueError:
        return None


def _lead_ok(po, ship):
    return po <= ship <= po + timedelta(days=LEAD_DAYS)


def _with_ship_year(po, ship):
    """po moved to the shipment's year (or the year before) when that puts it shortly before shipment."""
    if ship is None:
        return None
    for year in (ship.year, ship.year - 1):
        try:
            fixed = po.replace(year=year)
        except ValueError:  # 29 February
            continue
        if _lead_ok(fixed, ship):
            return fixed
    return None


def _check(row, today):
    """(problems, suggested PO Received Date, suggested Shipment Date) for one row; no problems if it looks right.

    Factories type dates day first; on a month-first PC Excel stores 06/12 as 12 June and keeps 18/12 as text.
    """
    problems, dates, valid = [], {}, {}
    for col in ("shipment_date", "po_received_date"):
        label, raw = HEADERS[col], row["_file"][col]
        d = date.fromisoformat(row[col]) if row[col] else None
        left_out = "; the row is left out of reports" if col == "shipment_date" else ""
        if raw is None:
            if col == "shipment_date":
                problems.append(f"{label} is empty{left_out}")
        elif d is None:
            problems.append(f"{label} is not a date{left_out}")
        else:
            if isinstance(raw, str):
                problems.append(f"{label} typed as text, read as day/month/year")
            if not FIRST_DATE <= d <= today:
                where = "in the future" if d > today else f"before {FIRST_DATE.year}"
                swapped = _swap(d)
                year_fixed = (_with_ship_year(d, dates["shipment_date"] if valid["shipment_date"] else None)
                              if col == "po_received_date" and d < FIRST_DATE else None)
                if swapped and FIRST_DATE <= swapped <= today:
                    problems.append(f"{label} is {where}; day and month look swapped")
                    d = swapped
                elif year_fixed:
                    problems.append(f"{label} is {where}; the year looks wrong")
                    d = year_fixed
                else:
                    problems.append(f"{label} is {where}")
        dates[col] = d
        valid[col] = d is not None and FIRST_DATE <= d <= today
    po, ship = dates["po_received_date"], dates["shipment_date"]
    if valid["po_received_date"] and valid["shipment_date"] and not _lead_ok(po, ship):
        po_swapped, ship_swapped = _swap(po), _swap(ship)
        if po_swapped and _lead_ok(po_swapped, ship):
            problems.append("PO Received Date: day and month look swapped")
            po = po_swapped
        elif po > ship:
            if ship_swapped and ship_swapped <= today and _lead_ok(po, ship_swapped):
                problems.append("Shipment Date: day and month look swapped")
                ship = ship_swapped
            else:
                problems.append("PO Received Date is after Shipment Date")
    return problems, po, ship


def check_dates(con, files, today=None):
    """Apply saved date corrections, then list the rows whose dates still look wrong.

    `files` is {file name: read_workbook()}. Rows with a saved correction get its dates in place and
    are not listed. Returns (rows corrected from saved corrections, issues). An issue is a dict of
    display fields, the suggested "PO Received Date" and "Shipment Date" (date or None), and "row":
    the row dict apply_dates() updates.
    """
    today = today or date.today()
    saved = {r["row_key"]: r for r in con.execute("SELECT row_key, po_received_date, shipment_date FROM date_fixes")}
    reused, issues = 0, []
    for name, data in files.items():
        for kind, rows in data.items():
            for r in rows:
                fix = saved.get(r["_key"])
                if fix:
                    r["po_received_date"], r["shipment_date"] = fix["po_received_date"], fix["shipment_date"]
                    reused += 1
                    continue
                problems, po, ship = _check(r, today)
                if problems:
                    issues.append({
                        "File": name, "Sheet": SHEETS[kind], "Excel row": r["_row"],
                        "Packaging supplier": r["packaging_supplier"], "eBiz Sales Order #": r["ebiz_so"],
                        "Carton QTY": r["carton_qty"],
                        "PO Received Date in file": _shown(r["_file"]["po_received_date"]),
                        "Shipment Date in file": _shown(r["_file"]["shipment_date"]),
                        "Problem": "; ".join(problems), "PO Received Date": po, "Shipment Date": ship, "row": r})
    return reused, issues


def apply_dates(issues, dates):
    """Give each issue's row its reviewed (PO Received Date, Shipment Date); save_many() remembers them."""
    for issue, (po, ship) in zip(issues, dates):
        row = issue["row"]
        row["po_received_date"] = po.isoformat() if po else None
        row["shipment_date"] = ship.isoformat() if ship else None
        row["_review"] = True


def saved_dates(con):
    """Saved date corrections, newest first. "key" is what forget_dates() takes."""
    def shown(iso):
        return _shown(date.fromisoformat(iso)) if iso else None
    return [{"Sheet": SHEETS[r["carton_type"]], "Packaging supplier": r["packaging_supplier"],
             "eBiz Sales Order #": r["ebiz_so"], "Carton QTY": r["carton_qty"],
             "PO Received Date in file": r["file_po"], "Shipment Date in file": r["file_ship"],
             "PO Received Date": shown(r["po_received_date"]), "Shipment Date": shown(r["shipment_date"]),
             "Saved at": r["saved_at"].replace("T", " ")[:16], "key": r["row_key"]}
            for r in con.execute("SELECT * FROM date_fixes ORDER BY saved_at DESC, rowid")]


def forget_dates(con, keys):
    """Delete saved corrections; those rows are checked again the next time they are imported."""
    with con:
        con.executemany("DELETE FROM date_fixes WHERE row_key = ?", [(k,) for k in keys])


def clashes(files):
    """Sheet + supplier + month groups found in more than one file. `files` is {file name: read_workbook()}."""
    seen = {}
    for name, data in files.items():
        for kind, rows in data.items():
            for supplier, month in {(r["supplier_key"], _month(r)) for r in rows}:
                seen.setdefault((kind, supplier, month), []).append(name)
    return [{"Sheet": SHEETS[kind], "Supplier": supplier, "Shipment month": month or "no valid date",
             "Files": " / ".join(names)}
            for (kind, supplier, month), names in sorted(seen.items(), key=str) if len(names) > 1]


def save(con, data, source_file):
    """Import one file. Returns the number of rows inserted per sheet."""
    return save_many(con, {source_file: data})[source_file]


def save_many(con, files):
    """Import several files in one transaction: all of them load, or none do.

    Per (supplier key, shipment month) in a file: delete those DB rows, insert the file's rows.
    Dates reviewed with apply_dates() are saved for the next import of the same rows.
    Raises ValueError if two files share a sheet + supplier + month. Returns {file: rows per sheet}.
    """
    clash = clashes(files)
    if clash:
        raise ValueError("Files overlap: " + "; ".join(
            f"{c['Sheet']} {c['Supplier']} {c['Shipment month']} in {c['Files']}" for c in clash))
    now = datetime.now().isoformat(timespec="seconds")
    cols = COLUMNS + ["supplier_key", "source_file", "imported_at"]
    with con:
        for name, data in files.items():
            for kind, rows in data.items():
                table = TABLES[kind]
                for key in {(r["supplier_key"], _month(r)) for r in rows}:
                    con.execute(f"DELETE FROM {table} WHERE {_SAME_KEY}", key)
                con.executemany(f"INSERT INTO {table} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                                [[r[c] for c in COLUMNS] + [r["supplier_key"], name, now] for r in rows])
                con.executemany("INSERT OR REPLACE INTO date_fixes VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                                [(r["_key"], kind, r["packaging_supplier"], r["ebiz_so"], r["carton_qty"],
                                  _shown(r["_file"]["po_received_date"]), _shown(r["_file"]["shipment_date"]),
                                  r["po_received_date"], r["shipment_date"], now)
                                 for r in rows if r.get("_review")])
    return {name: {kind: len(rows) for kind, rows in data.items()} for name, data in files.items()}


def _so(text):
    """An SO number as text, or None for blank and '0'."""
    text = (text or "").strip()
    return text if text and text != "0" else None


def _ebiz_value(col, value):
    if value is None:
        return None
    if col in EBIZ_REAL:
        try:
            return float(value)
        except (TypeError, ValueError):
            return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 7124561.0 -> 7124561
    return str(value).strip()


def read_ebiz(src):
    """Rows of one eBiz export (its first sheet) as dicts of EBIZ_COLUMNS. Raises ImportFileError.

    Order Number "A|B" gives so_1 "A" and so_2 "B"; vendor_key is supplier_key(vendor). Blank rows and rows
    with no Vendor (total lines) are skipped.
    """
    try:
        wb = openpyxl.load_workbook(src, data_only=True)
    except Exception as e:  # not an xlsx, corrupt or password protected
        raise ImportFileError(f"Can't open the workbook: {e}") from e
    ws = wb.worksheets[0]
    for n, header_row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        found = {_norm(h): i for i, h in enumerate(header_row)}
        if _norm(EBIZ_HEADERS["order_number"]) in found:
            break
    else:
        raise ImportFileError("eBiz header row not found (no 'Order Number' cell in rows 1-10).")
    index = {c: found.get(_norm(h)) for c, h in EBIZ_HEADERS.items()}
    for col in EBIZ_REQUIRED:
        if index[col] is None:
            raise ImportFileError(f"eBiz column '{EBIZ_HEADERS[col]}' not found in the header row (row {n}).")
    rows = []
    for values in ws.iter_rows(min_row=n + 1, values_only=True):
        row = {c: _ebiz_value(c, _clean(values[i])) if i is not None and i < len(values) else None
               for c, i in index.items()}
        if row["order_number"] is None and row["qty_shipped"] is None:
            continue
        row["vendor_key"] = supplier_key(row["vendor"])
        if row["vendor_key"] is None:  # a total/footer line: no vendor to report it under
            continue
        parts = (row["order_number"] or "").split("|")
        row["so_1"] = _so(parts[0])
        row["so_2"] = _so(parts[1]) if len(parts) > 1 else None
        rows.append(row)
    return rows


def ebiz_preview(files):
    """Per file and vendor key: spellings, rows, Qty Shipped and Total Rebate. `files` is {name: read_ebiz()}."""
    out = []
    for name, rows in files.items():
        groups = {}
        for r in rows:
            groups.setdefault(r["vendor_key"], []).append(r)
        for key, group in sorted(groups.items(), key=lambda g: str(g[0])):
            out.append({"File": name, "Vendor": key,
                        "Spellings in file": " / ".join(sorted({r["vendor"] for r in group} - {None})),
                        "Rows": len(group), "Qty Shipped": sum(r["qty_shipped"] or 0 for r in group),
                        "Total Rebate": round(sum(r["total_rebate"] or 0 for r in group), 2)})
    return out


def ebiz_clashes(files):
    """Vendor keys found in more than one eBiz file."""
    seen = {}
    for name, rows in files.items():
        for key in {r["vendor_key"] for r in rows}:
            seen.setdefault(key, []).append(name)
    return [{"Vendor": key, "Files": " / ".join(names)}
            for key, names in sorted(seen.items(), key=lambda s: str(s[0])) if len(names) > 1]


def replace_ebiz(con, files):
    """Delete every eBiz row, then insert the files' rows, in one transaction. Returns {file: rows inserted}.

    Raises ValueError, and changes nothing, when two files share a vendor key.
    """
    clash = ebiz_clashes(files)
    if clash:
        raise ValueError("Files overlap: " + "; ".join(f"{c['Vendor']} in {c['Files']}" for c in clash))
    now = datetime.now().isoformat(timespec="seconds")
    cols = EBIZ_COLUMNS + ["source_file", "imported_at"]
    with con:
        con.execute("DELETE FROM ebiz")
        for name, rows in files.items():
            con.executemany(f"INSERT INTO ebiz ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                            [[r[c] for c in EBIZ_COLUMNS] + [name, now] for r in rows])
    return {name: len(rows) for name, rows in files.items()}


def ebiz_status(con):
    """(row count, source file names, import time) of the current eBiz data, or None when there is none."""
    count, imported_at = con.execute("SELECT COUNT(*), MAX(imported_at) FROM ebiz").fetchone()
    if not count:
        return None
    files = [r[0] for r in con.execute("SELECT DISTINCT source_file FROM ebiz ORDER BY 1")]
    return count, files, imported_at


def ebiz_vendors(con):
    """Vendor keys in the eBiz data, sorted."""
    return [r[0] for r in con.execute(
        "SELECT DISTINCT vendor_key FROM ebiz WHERE vendor_key IS NOT NULL ORDER BY 1")]


def so_order(so):
    """Sort key for SO numbers: numeric ones by value, then text, then None."""
    if so is None:
        return (2, 0, "")
    return (0, int(so), "") if so.isdecimal() else (1, 0, so)


def cross_reference(con, supplier):
    """Tracker vs eBiz for one supplier key over all its tracker rows, any date (port of the old macros).

    Returns a dict:
    - "orders": one dict per tracker SO#, sorted by SO#: so, tracker_qty, tracker_rebate (PACD + Matalan
      rebate), ebiz_qty, ebiz_rebate, via_so2 (an eBiz row matched it through SO_2) and status
      ("Not in eBiz", "Tracker > eBiz", "Tracker < eBiz" or "Match").
    - "rows": {"STD": [...], "CUS": [...]} every tracker row of the supplier (view columns plus "so", the
      SO# or None), sorted by SO#, shipment date, id.
    - "ebiz_only": the supplier's eBiz rows whose order is not a tracker SO#, with "key" (SO_1, else SO_2).
    """
    rows = {}
    for kind in TABLES:
        found = [dict(r) for r in con.execute(f"SELECT * FROM v_{kind.lower()} WHERE supplier_key IS ?",
                                              (supplier,))]
        for r in found:
            r["so"] = _so(r["ebiz_so"])
        rows[kind] = sorted(found, key=lambda r: (so_order(r["so"]), r["shipment_date"] or "", r["id"]))
    orders = {}
    for r in rows["STD"] + rows["CUS"]:
        if r["so"]:
            o = orders.setdefault(r["so"], {"so": r["so"], "tracker_qty": 0, "tracker_rebate": 0.0,
                                            "ebiz_qty": 0, "ebiz_rebate": 0.0, "via_so2": False, "_found": 0})
            o["tracker_qty"] += r["carton_qty"] or 0
            o["tracker_rebate"] += (r["pacd"] or 0) + (r["matalan_rebate"] or 0)
    ebiz_only = []
    for e in (dict(r) for r in con.execute("SELECT * FROM ebiz WHERE vendor_key IS ? ORDER BY id", (supplier,))):
        if e["so_1"] in orders:
            e["key"], via_so2 = e["so_1"], False
        elif e["so_2"] in orders:
            e["key"], via_so2 = e["so_2"], True
        else:
            e["key"] = e["so_1"] or e["so_2"]
            ebiz_only.append(e)
            continue
        o = orders[e["key"]]
        o["ebiz_qty"] += e["qty_shipped"] or 0
        o["ebiz_rebate"] += e["total_rebate"] or 0
        o["via_so2"] = o["via_so2"] or via_so2
        o["_found"] += 1
    for o in orders.values():
        found = o.pop("_found")
        o["status"] = ("Not in eBiz" if not found else "Tracker > eBiz" if o["tracker_qty"] > o["ebiz_qty"]
                       else "Tracker < eBiz" if o["tracker_qty"] < o["ebiz_qty"] else "Match")
    return {"orders": sorted(orders.values(), key=lambda o: so_order(o["so"])), "rows": rows,
            "ebiz_only": ebiz_only}


def get_rates(con):
    """All rates, STD first, oldest first. Percentages are fractions (0.075 = 7.5%)."""
    return [dict(r) for r in con.execute(
        "SELECT carton_type, effective_from, pacd_pct, matalan_pct FROM rates "
        "ORDER BY carton_type DESC, effective_from")]


def save_rates(con, rates):
    """Replace every rate.

    Raises ValueError, and saves nothing, if a row is invalid or a dated shipment would be left without a rate.
    """
    if not rates:
        raise ValueError("At least one rate is needed.")
    rows, seen = [], set()
    for i, r in enumerate(rates, 1):
        kind = r.get("carton_type")
        if kind not in TABLES:
            raise ValueError(f"Row {i}: carton type must be STD or CUS.")
        try:
            start = date.fromisoformat(str(r.get("effective_from"))).isoformat()
        except ValueError:
            raise ValueError(f"Row {i}: 'From' must be a date (YYYY-MM-DD).") from None
        pcts = (r.get("pacd_pct"), r.get("matalan_pct"))
        if not all(isinstance(p, (int, float)) and 0 <= p <= 1 for p in pcts):
            raise ValueError(f"Row {i}: percentages must be between 0 and 100.")
        if (kind, start) in seen:
            raise ValueError(f"Row {i}: duplicate {kind} rate for {start}.")
        seen.add((kind, start))
        rows.append((kind, start, float(pcts[0]), float(pcts[1])))
    with con:
        con.execute("DELETE FROM rates")
        con.executemany("INSERT INTO rates VALUES (?, ?, ?, ?)", rows)
        n, first = _unrated(con)
        if n:  # raising here rolls the transaction back
            raise ValueError(f"{n} shipment(s) from {first} on would have no rate. "
                             f"Add a rate that starts on or before {first}.")


QUARTER_MONTHS = {1: "Dec-Mar", 2: "Apr-Jun", 3: "Jul-Aug", 4: "Sep-Nov"}


def quarter_label(fiscal_year, quarter):
    return f"Q{quarter} {fiscal_year} ({QUARTER_MONTHS[quarter]})"


def _where(supplier=None, suppliers=None, fiscal_year=None, quarters=None):
    sql, params = ["shipment_date IS NOT NULL"], []
    if supplier is not None:
        sql.append("supplier_key = ?")
        params.append(supplier)
    if suppliers:
        sql.append(f"supplier_key IN ({', '.join('?' * len(suppliers))})")
        params += list(suppliers)
    if fiscal_year is not None:
        sql.append("fiscal_year = ?")
        params.append(fiscal_year)
    if quarters:
        sql.append(f"quarter IN ({', '.join('?' * len(quarters))})")
        params += list(quarters)
    return " AND ".join(sql), params


def _dated(**filters):
    """SQL source for the dated rows of both views, and its parameters."""
    where, params = _where(**filters)
    return f"(SELECT * FROM v_std WHERE {where} UNION ALL SELECT * FROM v_cus WHERE {where})", params * 2


def rows(con, kind, **filters):
    """Dated rows of v_std or v_cus, sorted by shipment date."""
    where, params = _where(**filters)
    return [dict(r) for r in con.execute(
        f"SELECT * FROM v_{kind.lower()} WHERE {where} ORDER BY shipment_date, id", params)]


def editable_rows(con, kind, supplier=None, limit=None, offset=0):
    """Raw tracker rows for the web editor, including undated rows.

    Computed report fields are deliberately excluded: rates, rebates, PACD,
    fiscal year and quarter are recalculated by the views after every edit.
    """
    if kind not in TABLES:
        raise ValueError("Carton type must be STD or CUS.")
    where, params = "", []
    if supplier is not None:
        where, params = "WHERE supplier_key IS ?", [supplier]
    paging = ""
    if limit is not None:
        paging = " LIMIT ? OFFSET ?"
        params += [int(limit), int(offset)]
    return [dict(r) for r in con.execute(
        f"SELECT id, {', '.join(COLUMNS)}, supplier_key, source_file, imported_at "
        f"FROM {TABLES[kind]} {where} "
        f"ORDER BY shipment_date IS NULL, shipment_date, id{paging}", params)]


def editable_row_count(con, kind, supplier=None):
    """Number of raw rows matching one tracker-editor selection."""
    if kind not in TABLES:
        raise ValueError("Carton type must be STD or CUS.")
    if supplier is None:
        return con.execute(f"SELECT COUNT(*) FROM {TABLES[kind]}").fetchone()[0]
    return con.execute(
        f"SELECT COUNT(*) FROM {TABLES[kind]} WHERE supplier_key IS ?",
        (supplier,)).fetchone()[0]


def _manual_date(value, row_number, label):
    """Validate a date entered in the tracker editor and return ISO text."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError:
        parsed = _text_date(text)
        if parsed:
            return parsed.isoformat()
    raise ValueError(
        f"Row {row_number}: {label} must be YYYY-MM-DD or DD/MM/YYYY.")


def update_tracker_rows(con, kind, edited):
    """Save rows from the tracker editor in one transaction.

    ``kind`` fixes the destination table and each mapping must contain an
    existing ``id`` plus the imported columns. Invalid input rolls the whole
    edit back, so a partially edited tracker is never persisted.
    """
    if kind not in TABLES:
        raise ValueError("Carton type must be STD or CUS.")
    table = TABLES[kind]
    prepared, seen = [], set()
    for number, item in enumerate(edited, 1):
        try:
            row_id = int(item.get("id"))
        except (TypeError, ValueError):
            raise ValueError(f"Row {number}: row ID is invalid.") from None
        if row_id in seen:
            raise ValueError(f"Row {number}: duplicate row ID {row_id}.")
        seen.add(row_id)
        if not con.execute(f"SELECT 1 FROM {table} WHERE id = ?", (row_id,)).fetchone():
            raise ValueError(f"Row {number}: row ID {row_id} no longer exists. Refresh the page.")

        values = []
        for col in COLUMNS:
            value = _clean(item.get(col))
            if col in DATE_COLS:
                value = _manual_date(value, number, HEADERS[col])
            elif col in INT_COLS:
                if value is None or value == "":
                    value = None
                else:
                    try:
                        value = int(round(float(value)))
                    except (TypeError, ValueError):
                        raise ValueError(
                            f"Row {number}: {HEADERS[col]} must be a number.") from None
            elif col in REAL_COLS:
                if value is None or value == "":
                    value = None
                else:
                    try:
                        value = float(value)
                    except (TypeError, ValueError):
                        raise ValueError(
                            f"Row {number}: {HEADERS[col]} must be a number.") from None
            else:
                value = _clean(value)
                if value is not None:
                    value = str(value)
            values.append(value)
        key = supplier_key(values[COLUMNS.index("packaging_supplier")])
        prepared.append(values + [key, row_id])

    assignments = ", ".join(f"{col} = ?" for col in COLUMNS)
    now = datetime.now().isoformat(timespec="seconds")
    with con:
        con.executemany(
            f"UPDATE {table} SET {assignments}, supplier_key = ?, imported_at = ? WHERE id = ?",
            [values[:-1] + [now, values[-1]] for values in prepared])
    return len(prepared)


def dashboard_stats(con):
    """Small, set-oriented summary used by the website dashboard."""
    src, params = _dated()
    row = con.execute(
        f"SELECT COUNT(*), COALESCE(SUM(carton_qty), 0), "
        f"COUNT(DISTINCT supplier_key), MIN(shipment_date), MAX(shipment_date), "
        f"SUM(pacd_pct IS NULL) FROM {src}", params).fetchone()
    undated = sum(con.execute(
        f"SELECT COUNT(*) FROM {table} WHERE shipment_date IS NULL").fetchone()[0]
        for table in TABLES.values())
    return {"rows": row[0], "cartons": row[1], "suppliers": row[2],
            "first_date": row[3], "last_date": row[4],
            "unrated": row[5], "undated": undated}


def monthly_cartons(con):
    """Carton quantities by shipment month and packaging supplier."""
    src, params = _dated()
    return [dict(r) for r in con.execute(
        f"SELECT substr(shipment_date, 1, 7) AS month, supplier_key, "
        f"SUM(COALESCE(carton_qty, 0)) AS cartons FROM {src} "
        "GROUP BY month, supplier_key ORDER BY month, supplier_key", params)]


def top_tracker_entities(con, column, limit=10):
    """Top garment suppliers or factories by cartons in dated tracker rows."""
    if column not in {"supplier_name", "factory_name"}:
        raise ValueError("Unsupported tracker ranking column")
    src, params = _dated()
    return [dict(r) for r in con.execute(
        f"SELECT {column} AS name, SUM(COALESCE(carton_qty, 0)) AS cartons "
        f"FROM {src} WHERE {column} IS NOT NULL AND TRIM({column}) <> '' "
        f"GROUP BY {column} ORDER BY cartons DESC, name LIMIT ?",
        [*params, int(limit)])]


def _entity_name(value):
    """Trim and collapse whitespace in a managed supplier/factory name."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def sync_assignment_catalog(con):
    """Discover tracker factories and seed packaging suppliers once.

    Factories continue to be discovered as tracker data changes. Packaging
    suppliers are only seeded once so a supplier deleted by an administrator
    is not recreated from historical tracker rows.
    """
    raw = " UNION ALL ".join(
        f"SELECT factory_name, packaging_supplier, supplier_key FROM {table}"
        for table in TABLES.values())
    rows = con.execute(
        f"SELECT factory_name, packaging_supplier, supplier_key FROM ({raw}) "
        "WHERE factory_name IS NOT NULL OR packaging_supplier IS NOT NULL").fetchall()
    now = datetime.now().isoformat(timespec="seconds")
    factories = sorted({_entity_name(r["factory_name"]) for r in rows
                        if _entity_name(r["factory_name"])}, key=str.casefold)
    with con:
        con.executemany(
            "INSERT OR IGNORE INTO assignment_factories (name, created_at) VALUES (?, ?)",
            [(name, now) for name in factories])
        seeded = con.execute(
            "SELECT value FROM app_meta WHERE key = 'assignment_suppliers_seeded'").fetchone()
        if seeded is None:
            counts = {}
            for row in rows:
                name = _entity_name(row["packaging_supplier"])
                key = row["supplier_key"] or supplier_key(name)
                if name and key:
                    counts[(key, name)] = counts.get((key, name), 0) + 1
            chosen = {}
            for (key, name), count in counts.items():
                current = chosen.get(key)
                if current is None or (-count, name.casefold()) < (-current[0], current[1].casefold()):
                    chosen[key] = (count, name)
            if chosen:
                con.executemany(
                    "INSERT OR IGNORE INTO packaging_suppliers (name, created_at) VALUES (?, ?)",
                    [(item[1], now) for item in sorted(chosen.values(), key=lambda x: x[1].casefold())])
                con.execute(
                    "INSERT INTO app_meta (key, value) VALUES ('assignment_suppliers_seeded', '1')")


def list_packaging_suppliers(con):
    """Managed packaging suppliers and their factory assignment counts."""
    return [dict(row) for row in con.execute(
        "SELECT ps.id, ps.name, COUNT(fa.factory_id) AS assignment_count "
        "FROM packaging_suppliers ps LEFT JOIN factory_assignments fa "
        "ON fa.packaging_supplier_id = ps.id "
        "GROUP BY ps.id, ps.name ORDER BY ps.name COLLATE NOCASE")]


def add_packaging_supplier(con, name):
    name = _entity_name(name)
    if not name:
        raise ValueError("Enter a packaging supplier name.")
    try:
        with con:
            cur = con.execute(
                "INSERT INTO packaging_suppliers (name, created_at) VALUES (?, ?)",
                (name, datetime.now().isoformat(timespec="seconds")))
            con.execute(
                "INSERT OR IGNORE INTO app_meta (key, value) "
                "VALUES ('assignment_suppliers_seeded', '1')")
    except sqlite3.IntegrityError:
        raise ValueError(f"A packaging supplier named {name} already exists.") from None
    return cur.lastrowid


def rename_packaging_supplier(con, supplier_id, name):
    name = _entity_name(name)
    if not name:
        raise ValueError("Enter a packaging supplier name.")
    if con.execute("SELECT 1 FROM packaging_suppliers WHERE id = ?", (supplier_id,)).fetchone() is None:
        raise ValueError("Packaging supplier not found.")
    try:
        with con:
            con.execute("UPDATE packaging_suppliers SET name = ? WHERE id = ?", (name, supplier_id))
    except sqlite3.IntegrityError:
        raise ValueError(f"A packaging supplier named {name} already exists.") from None


def delete_packaging_supplier(con, supplier_id):
    row = con.execute(
        "SELECT ps.name, COUNT(fa.factory_id) AS factory_count "
        "FROM packaging_suppliers ps LEFT JOIN factory_assignments fa "
        "ON fa.packaging_supplier_id = ps.id WHERE ps.id = ? GROUP BY ps.id",
        (supplier_id,)).fetchone()
    if row is None:
        raise ValueError("Packaging supplier not found.")
    if row["factory_count"]:
        raise ValueError(
            f"{row['name']} can't be deleted while assigned to "
            f"{row['factory_count']:,} factories. Remove those assignments first.")
    with con:
        con.execute("DELETE FROM packaging_suppliers WHERE id = ?", (supplier_id,))


def list_assignment_factories(con):
    """Tracker-discovered factories with priority-ordered assignments."""
    factories = [dict(row) for row in con.execute(
        "SELECT id, name FROM assignment_factories ORDER BY name COLLATE NOCASE")]
    assigned = {}
    for row in con.execute(
            "SELECT fa.factory_id, ps.id, ps.name FROM factory_assignments fa "
            "JOIN packaging_suppliers ps ON ps.id = fa.packaging_supplier_id "
            "ORDER BY fa.factory_id, fa.priority"):
        assigned.setdefault(row["factory_id"], []).append((row["id"], row["name"]))
    for factory in factories:
        pairs = assigned.get(factory["id"], [])
        factory["assigned_ids"] = [item[0] for item in pairs]
        factory["assigned_names"] = [item[1] for item in pairs]
    return factories


def set_factory_assignments(con, factory_id, supplier_ids):
    """Replace one factory's supplier set, preserving selection order."""
    supplier_ids = list(dict.fromkeys(int(value) for value in supplier_ids))
    if con.execute("SELECT 1 FROM assignment_factories WHERE id = ?", (factory_id,)).fetchone() is None:
        raise ValueError("Factory not found.")
    if supplier_ids:
        marks = ", ".join("?" * len(supplier_ids))
        found = {row[0] for row in con.execute(
            f"SELECT id FROM packaging_suppliers WHERE id IN ({marks})", supplier_ids)}
        if found != set(supplier_ids):
            raise ValueError("One or more packaging suppliers no longer exist.")
    with con:
        con.execute("DELETE FROM factory_assignments WHERE factory_id = ?", (factory_id,))
        con.executemany(
            "INSERT INTO factory_assignments "
            "(factory_id, packaging_supplier_id, priority) VALUES (?, ?, ?)",
            [(factory_id, supplier_id, priority)
             for priority, supplier_id in enumerate(supplier_ids, start=1)])


def export_factory_assignments(con):
    """Current factory assignments as an import-compatible XLSX workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Factory Assignments"
    ws.append(["Factory", "Packaging Suppliers"])
    for factory in list_assignment_factories(con):
        ws.append([factory["name"], " | ".join(factory["assigned_names"])])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 68
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def import_factory_assignments(con, content):
    """Apply an XLSX exported by export_factory_assignments atomically."""
    try:
        ws = openpyxl.load_workbook(BytesIO(content), data_only=True).active
    except Exception as exc:
        raise ValueError(f"Can't open the assignments workbook: {exc}") from exc
    try:
        header_values = next(ws.iter_rows(values_only=True))
    except StopIteration:
        raise ValueError("Assignments workbook is empty.") from None
    headers = {_norm(value): index for index, value in enumerate(header_values)}
    factory_col = headers.get("factory")
    suppliers_col = headers.get("packaging suppliers")
    if factory_col is None or suppliers_col is None:
        raise ValueError("Assignments workbook requires Factory and Packaging Suppliers columns.")
    factories = {row["name"].casefold(): row for row in list_assignment_factories(con)}
    suppliers = {row["name"].casefold(): row for row in list_packaging_suppliers(con)}
    prepared, errors, seen = [], [], set()
    for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        factory_name = _entity_name(values[factory_col] if factory_col < len(values) else None)
        raw_suppliers = _entity_name(values[suppliers_col] if suppliers_col < len(values) else None)
        if not factory_name and not raw_suppliers:
            continue
        factory = factories.get(factory_name.casefold())
        if factory is None:
            errors.append(f"Row {number}: unknown factory {factory_name or '(blank)' }.")
            continue
        if factory["id"] in seen:
            errors.append(f"Row {number}: factory {factory_name} appears more than once.")
            continue
        seen.add(factory["id"])
        names = [_entity_name(value) for value in re.split(r"\s*[|;,]\s*", raw_suppliers) if _entity_name(value)]
        missing = [name for name in names if name.casefold() not in suppliers]
        if missing:
            errors.append(f"Row {number}: unknown packaging supplier(s): {', '.join(missing)}.")
            continue
        prepared.append((factory["id"], [suppliers[name.casefold()]["id"] for name in names]))
    if errors:
        raise ValueError("\n".join(errors))
    with con:
        for factory_id, supplier_ids in prepared:
            con.execute("DELETE FROM factory_assignments WHERE factory_id = ?", (factory_id,))
            con.executemany(
                "INSERT INTO factory_assignments "
                "(factory_id, packaging_supplier_id, priority) VALUES (?, ?, ?)",
                [(factory_id, supplier_id, priority)
                 for priority, supplier_id in enumerate(supplier_ids, start=1)])
    return len(prepared)


def suppliers(con):
    """Supplier keys of the dated rows, sorted."""
    src, params = _dated()
    return [r[0] for r in con.execute(
        f"SELECT DISTINCT supplier_key FROM {src} WHERE supplier_key IS NOT NULL ORDER BY 1", params)]


def tracker_suppliers(con):
    """Supplier keys from every raw tracker row, including undated rows."""
    sql = " UNION ".join(
        f"SELECT supplier_key FROM {table} WHERE supplier_key IS NOT NULL"
        for table in TABLES.values())
    return [r[0] for r in con.execute(f"SELECT supplier_key FROM ({sql}) ORDER BY 1")]


def supplier_name(con, key):
    """Most common packaging supplier spelling of a supplier key (alphabetical first on a tie), else the key."""
    src, params = _dated(supplier=key)
    row = con.execute(f"SELECT packaging_supplier FROM {src} WHERE packaging_supplier IS NOT NULL "
                      f"GROUP BY packaging_supplier ORDER BY COUNT(*) DESC, packaging_supplier LIMIT 1",
                      params).fetchone()
    return row[0] if row else key


def fiscal_years(con):
    src, params = _dated()
    return [r[0] for r in con.execute(f"SELECT DISTINCT fiscal_year FROM {src} ORDER BY 1", params)]


def quarters(con, **filters):
    src, params = _dated(**filters)
    return [r[0] for r in con.execute(f"SELECT DISTINCT quarter FROM {src} ORDER BY 1", params)]


def date_range(con, **filters):
    """(first, last) shipment date of the matching rows, or None."""
    src, params = _dated(**filters)
    first, last = con.execute(f"SELECT MIN(shipment_date), MAX(shipment_date) FROM {src}", params).fetchone()
    return (date.fromisoformat(first), date.fromisoformat(last)) if first else None


def _unrated(con):
    """(count, first shipment date) of the dated rows that have no rate."""
    src, params = _dated()
    return con.execute(f"SELECT COUNT(*), MIN(shipment_date) FROM {src} WHERE pacd_pct IS NULL", params).fetchone()


def unrated_count(con):
    """Dated rows with no rate: shipped before the first rate of their carton type."""
    return _unrated(con)[0]


def quarter_totals(con, kind, supplier, fiscal_year):
    """{quarter: (qty, total amount, pacd, matalan rebate)} for one supplier and fiscal year."""
    where, params = _where(supplier=supplier, fiscal_year=fiscal_year)
    return {r[0]: tuple(r[1:]) for r in con.execute(
        f"SELECT quarter, SUM(carton_qty), SUM(total_amount), SUM(pacd), SUM(matalan_rebate) "
        f"FROM v_{kind.lower()} WHERE {where} GROUP BY quarter", params)}


def quarter_rates(con, kind, supplier, fiscal_year):
    """{quarter: (pacd %, matalan %)} for one supplier and fiscal year; None where the rows don't share one rate."""
    where, params = _where(supplier=supplier, fiscal_year=fiscal_year)
    out = {}
    for q, n_pacd, pacd, n_matalan, matalan, unrated in con.execute(
            f"SELECT quarter, COUNT(DISTINCT pacd_pct), MIN(pacd_pct), COUNT(DISTINCT matalan_pct), "
            f"MIN(matalan_pct), SUM(pacd_pct IS NULL) FROM v_{kind.lower()} WHERE {where} GROUP BY quarter", params):
        out[q] = (pacd if n_pacd == 1 and not unrated else None, matalan if n_matalan == 1 and not unrated else None)
    return out


def rate_headers(rows):
    """Headers for the PACD and rebate columns; they show the % when every row has the same rate."""
    def label(key, after_pct, plain):
        pcts = {r[key] for r in rows}
        return f"{round(pcts.pop() * 100, 4):g}% {after_pct}" if len(pcts) == 1 and None not in pcts else plain
    return label("pacd_pct", "for PACD", "PACD"), label("matalan_pct", "Rebates for Matalan", "Rebates for Matalan")
